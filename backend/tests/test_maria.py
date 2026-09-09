"""
Testes unitários para os módulos do projeto MARIA.
Cobre lógica determinística de chat_session.py e tools_schema.py.

Executar:
    python -m unittest test_maria.py
"""

import unittest
import tempfile
import os
import json
import time
import pandas as pd
from unittest.mock import patch, MagicMock
from backend.core.chat_session import ChatSession, interpretar_confirmacao
from backend.core.llama_client import _montar_mensagens_com_reforco
from backend.core.tools_schema import (
    simular_execucao_ferramenta,
    executar_ferramenta_real,
    validar_argumentos_obrigatorios,
    FERRAMENTA_CRIAR_PLANILHA,
    FERRAMENTA_CRIAR_DOCUMENTO,
    FERRAMENTA_EDITAR_PLANILHA,
    executar_ferramenta_leitura,
    FERRAMENTAS_LEITURA,
)
from backend.core.excel_handler import criar_planilha_real, editar_planilha_real
from backend.core.file_utils import (
    gerar_nome_unico,
    garantir_pasta_arquivos,
    resolver_caminho_permitido,
    listar_arquivos,
    ler_documento,
)
from backend.core.session_storage import salvar_sessao, listar_sessoes_salvas, carregar_sessao
from backend.benchmarks.maria_bench.analysis.language_check import resposta_em_portugues
from backend.benchmarks.maria_bench.analysis.metrics import calculate_maria_metrics
from backend.benchmarks.maria_bench.tasks.task_schema import MariaTaskResult


class TestChatSession(unittest.TestCase):
    """Testes para a classe ChatSession."""
    
    def setUp(self):
        """Configura uma sessão fresca para cada teste."""
        self.sessao = ChatSession(max_mensagens=5)
    
    def test_adicionar_mensagem_user(self):
        """Testa adição de mensagem do usuário."""
        self.sessao.adicionar_mensagem("user", "Olá")
        self.assertEqual(self.sessao.contar_mensagens(), 1)
        self.assertEqual(self.sessao.get_ultima_mensagem_usuario(), "Olá")
    
    def test_adicionar_mensagem_assistant(self):
        """Testa adição de mensagem do assistente."""
        self.sessao.adicionar_mensagem("assistant", "Olá! Como posso ajudar?")
        self.assertEqual(self.sessao.contar_mensagens(), 1)
    
    def test_limite_historico_preserva_system_prompt(self):
        """Testa que o limite de histórico funciona corretamente sem system prompt interno."""
        # O histórico interno não contém mais system prompt
        # Adicionar 10 mensagens (mais que o limite de 5)
        for i in range(10):
            self.sessao.adicionar_mensagem("user", f"Mensagem {i}")
            self.sessao.adicionar_mensagem("assistant", f"Resposta {i}")
        
        # Verificar que get_historico_com_system injeta exatamente UM system prompt
        historico_completo = self.sessao.get_historico_com_system()
        self.assertEqual(historico_completo[0]["role"], "system")
        self.assertIn("MARIA", historico_completo[0]["content"])
        
        # Verificar que há apenas UM system prompt no histórico retornado
        system_prompts = [m for m in historico_completo if m["role"] == "system"]
        self.assertEqual(len(system_prompts), 1)
        
        # Verificar que histórico tem no máximo 5 mensagens user/assistant + 1 system
        self.assertLessEqual(len(historico_completo), 6)  # 5 msgs + 1 system
    
    def test_get_historico_com_system_adiciona_automaticamente(self):
        """Testa que get_historico_com_system adiciona system prompt se não existir."""
        self.sessao.adicionar_mensagem("user", "Olá")
        
        historico = self.sessao.get_historico_com_system()
        
        self.assertEqual(len(historico), 2)  # system + user msg
        self.assertEqual(historico[0]["role"], "system")
        self.assertIn("MARIA", historico[0]["content"])

    def test_system_prompt_exige_portugues(self):
        """Testa que o prompt exige português do Brasil em qualquer resposta."""
        self.assertIn("português do Brasil", ChatSession.SYSTEM_PROMPT)
    
    def test_get_historico_sem_system_exclui_system_prompt(self):
        """Testa que get_historico_sem_system exclui system prompt."""
        self.sessao.adicionar_mensagem("user", "Olá")
        
        historico_com_system = self.sessao.get_historico_com_system()
        historico_sem_system = self.sessao.get_historico_sem_system()
        
        self.assertEqual(len(historico_com_system), 2)  # system + user msg
        self.assertEqual(len(historico_sem_system), 1)  # apenas user msg
        
        # Nenhum item em sem_system deve ter role system
        for msg in historico_sem_system:
            self.assertNotEqual(msg["role"], "system")
    
    def test_limpar_historico(self):
        """Testa que limpar_historico esvazia o histórico interno."""
        self.sessao.adicionar_mensagem("user", "Olá")
        self.sessao.limpar_historico()
        
        self.assertEqual(self.sessao.contar_mensagens(), 0)
        self.assertEqual(self.sessao.historico, [])
        historico = self.sessao.get_historico_com_system()
        self.assertEqual(len(historico), 1)
        self.assertEqual(historico[0]["role"], "system")
    
    def test_role_invalido_raises_error(self):
        """Testa que role inválido levanta ValueError."""
        with self.assertRaises(ValueError):
            self.sessao.adicionar_mensagem("invalid_role", "teste")
    
    def test_acao_pendente_inicialmente_vazia(self):
        """Testa que ação pendente inicia como None."""
        self.assertIsNone(self.sessao.acao_pendente)
        self.assertFalse(self.sessao.tem_acao_pendente())
    
    def test_definir_e_limpar_acao_pendente(self):
        """Testa definição e limpeza de ação pendente."""
        tool_call = {"name": "criar_planilha", "arguments": {"nome_arquivo": "teste"}}
        self.sessao.definir_acao_pendente(tool_call)
        
        self.assertTrue(self.sessao.tem_acao_pendente())
        self.assertEqual(self.sessao.acao_pendente["name"], "criar_planilha")
        self.assertEqual(self.sessao.acao_pendente["arguments"]["nome_arquivo"], "teste")
        self.assertEqual(self.sessao.tentativas_confirmacao_ambigua, 0)
        
        self.sessao.limpar_acao_pendente()
        self.assertFalse(self.sessao.tem_acao_pendente())
        self.assertIsNone(self.sessao.acao_pendente)


class TestInterpretarConfirmacao(unittest.TestCase):
    """Testes para a função interpretar_confirmacao."""
    
    def test_respostas_afirmativas(self):
        """Testa respostas afirmativas retornando True."""
        afirmativas = ["sim", "Sim", "SIM", "pode", "confirmo", "ok", "OK", 
                       "vai", "isso", "claro", "com certeza", "certeza", "vale", "bora"]
        for resposta in afirmativas:
            with self.subTest(resposta=resposta):
                self.assertIs(interpretar_confirmacao(resposta), True)
    
    def test_respostas_negativas(self):
        """Testa respostas negativas retornando False."""
        negativas = ["não", "Não", "NAO", "nao", "cancela", "para", "esquece",
                     "jamais", "de jeito nenhum", "nem pensar", "aborta", "desiste"]
        for resposta in negativas:
            with self.subTest(resposta=resposta):
                self.assertIs(interpretar_confirmacao(resposta), False)
    
    def test_respostas_ambiguas(self):
        """Testa respostas ambíguas retornando None."""
        ambiguas = ["talvez", "hummm", "bla bla", "", "   ", "quero ver", "depois"]
        for resposta in ambiguas:
            with self.subTest(resposta=resposta):
                self.assertIsNone(interpretar_confirmacao(resposta))


class TestExecucaoReal(unittest.TestCase):
    """Testes para execução real de criação de arquivos."""
    
    def setUp(self):
        """Configura pasta temporária para testes."""
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_pasta = os.environ.get("PASTA_ARQUIVOS_GERADOS")
        os.environ["PASTA_ARQUIVOS_GERADOS"] = self.temp_dir.name
    
    def tearDown(self):
        """Limpa pasta temporária."""
        self.temp_dir.cleanup()
        if self.original_pasta:
            os.environ["PASTA_ARQUIVOS_GERADOS"] = self.original_pasta
        else:
            os.environ.pop("PASTA_ARQUIVOS_GERADOS", None)
    
    def test_criar_planilha_real(self):
        """Testa criação real de planilha Excel."""
        argumentos = {
            "nome_arquivo": "teste_planilha",
            "colunas": ["Data", "Descrição", "Valor"],
            "descricao": "Planilha de teste"
        }
        
        resultado = executar_ferramenta_real("criar_planilha", argumentos)
        
        self.assertIn("Planilha criada com sucesso:", resultado)
        caminho = resultado.split(": ")[1]
        self.assertTrue(os.path.exists(caminho))
        self.assertTrue(caminho.endswith(".xlsx"))
    
    def test_criar_documento_real(self):
        """Testa criação real de documento Word."""
        argumentos = {
            "nome_arquivo": "teste_documento",
            "titulo": "Documento de Teste",
            "conteudo": "Primeiro parágrafo do documento.\n\nSegundo parágrafo do documento."
        }
        
        resultado = executar_ferramenta_real("criar_documento", argumentos)
        
        self.assertIn("Documento criado com sucesso:", resultado)
        caminho = resultado.split(": ")[1]
        self.assertTrue(os.path.exists(caminho))
        self.assertTrue(caminho.endswith(".docx"))

    def test_editar_planilha_real_sobrescreve_arquivo_existente(self):
        """Testa que editar_planilha_real sobrescreve uma planilha existente."""
        criar_planilha_real("teste_edicao", ["A", "B"], "original")

        caminho = editar_planilha_real(
            "teste_edicao",
            colunas=["X", "Y", "Z"],
            linhas=[{"X": 1, "Y": 2, "Z": 3}]
        )

        self.assertTrue(os.path.exists(caminho))
        self.assertTrue(caminho.endswith("teste_edicao.xlsx"))


class TestBenchmarkLanguageCompliance(unittest.TestCase):
    """Testes para a checagem de conformidade de idioma do benchmark."""

    def test_resposta_em_portugues_identifica_portugues(self):
        self.assertTrue(resposta_em_portugues("Por favor, responda em português."))
        self.assertTrue(resposta_em_portugues("Claro, vou ajudar você com isso."))

    def test_resposta_em_portugues_rejeita_inglês_significativo(self):
        self.assertFalse(resposta_em_portugues("Please provide the details in English."))
        self.assertFalse(resposta_em_portugues("Sure, I can help you with that."))

    def test_resposta_em_portugues_tolerancia_a_texto_misto(self):
        self.assertFalse(resposta_em_portugues("Por favor, please continue em português."))
        self.assertTrue(resposta_em_portugues("Por favor, não use termos em inglês."))


class TestBenchmarkMetrics(unittest.TestCase):
    """Testes para cálculo de métricas agregadas do benchmark."""

    def test_calculate_maria_metrics_includes_language_compliance(self):
        results = [
            MariaTaskResult(
                task_id=1,
                task_name="Teste 1",
                category="categoria",
                model="maria",
                tool_detected="criar_planilha",
                tool_correct=True,
                confirmation_completed=True,
                keyword_match=True,
                runtime_ok=True,
                final_message="Tudo certo.",
                latency_ms=100.0,
                errors=[],
                raw_tool_args={},
                language_ok=True,
            ),
            MariaTaskResult(
                task_id=2,
                task_name="Teste 2",
                category="categoria",
                model="maria",
                tool_detected="criar_planilha",
                tool_correct=False,
                confirmation_completed=False,
                keyword_match=False,
                runtime_ok=False,
                final_message="Houve um problema.",
                latency_ms=200.0,
                errors=[{"kind": "LlamaClientError", "message": "Falha"}],
                raw_tool_args={},
                language_ok=False,
            ),
        ]

        metrics = calculate_maria_metrics(results)

        self.assertEqual(metrics.total_tasks, 2)
        self.assertAlmostEqual(metrics.language_compliance_rate, 0.5)
        self.assertEqual(metrics.error_distribution.get("LlamaClientError"), 1)

    def test_calculate_maria_metrics_includes_avg_tokens_per_second(self):
        results = [
            MariaTaskResult(
                task_id=1,
                task_name="Teste 1",
                category="categoria",
                model="maria",
                tool_detected="criar_planilha",
                tool_correct=True,
                confirmation_completed=True,
                keyword_match=True,
                runtime_ok=True,
                final_message="Tudo certo.",
                latency_ms=100.0,
                errors=[],
                raw_tool_args={},
                language_ok=True,
                tokens_gerados=120,
                tokens_por_segundo=20.0,
            ),
            MariaTaskResult(
                task_id=2,
                task_name="Teste 2",
                category="categoria",
                model="maria",
                tool_detected="criar_planilha",
                tool_correct=False,
                confirmation_completed=False,
                keyword_match=False,
                runtime_ok=False,
                final_message="Erro.",
                latency_ms=200.0,
                errors=[{"kind": "LlamaClientError", "message": "Falha"}],
                raw_tool_args={},
                language_ok=False,
                tokens_gerados=80,
                tokens_por_segundo=10.0,
            ),
        ]

        metrics = calculate_maria_metrics(results)
        self.assertAlmostEqual(metrics.avg_tokens_por_segundo, 15.0)

    def test_resolver_tool_call_final_aceita_json_plano(self):
        """O streaming do modelo devolve tool calls no formato JSON plano."""
        from backend.core.llama_client import LlamaClient

        cliente = LlamaClient()
        tool_call, fonte, nome_bruto, fallbacks = cliente._resolver_tool_call_final(
            tc_detectada_via_delta=False,
            tc_nome_acumulado="",
            tc_args_acumulado="",
            conteudo_acumulado='{"ferramenta": "criar_planilha", "nome_arquivo": "gastos", "colunas": ["Data", "Valor"]}',
        )

        self.assertIsNotNone(tool_call)
        self.assertEqual(fonte, "json")
        self.assertEqual(fallbacks, [])
        self.assertEqual(tool_call["name"], "criar_planilha")
        self.assertEqual(tool_call["arguments"]["nome_arquivo"], "gastos")
        self.assertEqual(tool_call["arguments"]["colunas"], ["Data", "Valor"])

    def test_diagnostico_falha_sem_erro_detecta_tool_incorreto(self):
        result = MariaTaskResult(
            task_id=3,
            task_name="Teste 3",
            category="categoria",
            model="maria",
            tool_detected="editar_planilha",
            tool_correct=False,
            confirmation_completed=True,
            keyword_match=True,
            runtime_ok=True,
            final_message="Ferramenta incorreta.",
            latency_ms=150.0,
            errors=[],
            raw_tool_args={},
            language_ok=True,
        )
        from backend.benchmarks.maria_bench.analysis.report import _diagnosticar_falha
        self.assertEqual(_diagnosticar_falha(result), "Tool call incorreto ou ferramenta inesperada")

    def test_diagnostico_falha_por_idioma_incorreto(self):
        result = MariaTaskResult(
            task_id=4,
            task_name="Teste 4",
            category="categoria",
            model="maria",
            tool_detected="criar_documento",
            tool_correct=True,
            confirmation_completed=True,
            keyword_match=True,
            runtime_ok=True,
            final_message="This is a response in English.",
            latency_ms=120.0,
            errors=[],
            raw_tool_args={},
            language_ok=False,
        )
        from backend.benchmarks.maria_bench.analysis.report import _diagnosticar_falha
        self.assertEqual(_diagnosticar_falha(result), "Resposta em idioma incorreto")

    def test_editar_planilha_real_arquivo_inexistente_levanta_value_error(self):
        """Testa erro ao editar uma planilha inexistente."""
        with self.assertRaises(ValueError):
            editar_planilha_real("planilha_que_nao_existe", colunas=["A", "B"])
    def test_confirmation_rate_elegiveis_exclui_tasks_sem_confirmacao(self):
        """A taxa sobre elegíveis mede só tasks com confirm_sequence —
        elimina o efeito cascata de parser em tarefas sem confirmação."""
        elegivel_ok = MariaTaskResult(
            task_id=1, task_name="T1", category="c", model="m",
            tool_detected="criar_planilha", tool_correct=True,
            confirmation_completed=True, keyword_match=True, runtime_ok=True,
            final_message="ok", latency_ms=1.0, errors=[], raw_tool_args={},
            confirmacao_elegivel=True,
        )
        elegivel_falha = MariaTaskResult(
            task_id=2, task_name="T2", category="c", model="m",
            tool_detected=None, tool_correct=False,
            confirmation_completed=False, keyword_match=False, runtime_ok=True,
            final_message="falha", latency_ms=1.0, errors=[], raw_tool_args={},
            confirmacao_elegivel=True,
        )
        nao_elegivel = MariaTaskResult(
            task_id=3, task_name="T3", category="c", model="m",
            tool_detected=None, tool_correct=True,
            confirmation_completed=True, keyword_match=True, runtime_ok=True,
            final_message="ok", latency_ms=1.0, errors=[], raw_tool_args={},
            confirmacao_elegivel=False,
        )

        metrics = calculate_maria_metrics([elegivel_ok, elegivel_falha, nao_elegivel])

        # Taxa global inclui a não elegível (2/3); a elegível ignora (1/2).
        self.assertAlmostEqual(metrics.confirmation_success_rate, 2 / 3)
        self.assertAlmostEqual(metrics.confirmation_success_rate_elegiveis, 0.5)

    def test_confirmation_rate_elegiveis_none_sem_elegiveis(self):
        sem_elegivel = MariaTaskResult(
            task_id=1, task_name="T1", category="c", model="m",
            tool_detected=None, tool_correct=True,
            confirmation_completed=True, keyword_match=True, runtime_ok=True,
            final_message="ok", latency_ms=1.0, errors=[], raw_tool_args={},
            confirmacao_elegivel=False,
        )
        metrics = calculate_maria_metrics([sem_elegivel])
        self.assertIsNone(metrics.confirmation_success_rate_elegiveis)

    def test_parse_suspeito_contabilizado(self):
        suspeito = MariaTaskResult(
            task_id=1, task_name="T1", category="c", model="m",
            tool_detected=None, tool_correct=False,
            confirmation_completed=False, keyword_match=False, runtime_ok=True,
            final_message="falha", latency_ms=1.0, errors=[], raw_tool_args={},
            parse_suspeito=True,
        )
        limpo = MariaTaskResult(
            task_id=2, task_name="T2", category="c", model="m",
            tool_detected=None, tool_correct=True,
            confirmation_completed=True, keyword_match=True, runtime_ok=True,
            final_message="ok", latency_ms=1.0, errors=[], raw_tool_args={},
            parse_suspeito=False,
        )
        metrics = calculate_maria_metrics([suspeito, limpo])
        self.assertEqual(metrics.parse_suspeito_count, 1)


class TestValidacaoArgumentos(unittest.TestCase):
    """Testes para validação de campos obrigatórios em ferramentas."""

    def test_validar_argumentos_obrigatorios_ausente_levanta_value_error(self):
        """Campo obrigatório ausente deve ser rejeitado antes da execução."""
        with self.assertRaisesRegex(ValueError, "colunas"):
            validar_argumentos_obrigatorios(
                "criar_planilha",
                {"nome_arquivo": "teste"}
            )

    def test_validar_argumentos_obrigatorios_vazio_levanta_value_error(self):
        """String vazia deve ser tratada como campo obrigatório ausente."""
        with self.assertRaisesRegex(ValueError, "titulo"):
            validar_argumentos_obrigatorios(
                "criar_documento",
                {"nome_arquivo": "x", "titulo": "   ", "conteudo": "y"}
            )

    def test_validar_argumentos_obrigatorios_colunas_como_string_levanta_value_error(self):
        """'colunas' como string única deve ser rejeitado pelo validador (tipo)."""
        with self.assertRaisesRegex(ValueError, "lista de strings"):
            validar_argumentos_obrigatorios(
                "criar_planilha",
                {"nome_arquivo": "gastos", "colunas": "Data, Valor"}
            )

    def test_validar_argumentos_obrigatorios_nome_arquivo_path_traversal_levanta_value_error(self):
        """nome_arquivo com path traversal deve ser rejeitado pela sanitização."""
        with self.assertRaisesRegex(ValueError, "path traversal"):
            validar_argumentos_obrigatorios(
                "criar_planilha",
                {"nome_arquivo": "../../teste_seguro", "colunas": ["Data"]}
            )

    def test_validar_argumentos_obrigatorios_tool_call_valida_nao_levanta_excecao(self):
        """Tool call de escrita válida não deve levantar exceção."""
        validar_argumentos_obrigatorios(
            "criar_planilha",
            {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]}
        )
        validar_argumentos_obrigatorios(
            "criar_documento",
            {"nome_arquivo": "relatorio", "titulo": "Relatório", "conteudo": "Texto"}
        )

    def test_validar_argumentos_obrigatorios_colunas_com_item_dict_levanta_value_error(self):
        """Item de 'colunas' que não é string deve ser rejeitado (validação por item)."""
        with self.assertRaisesRegex(ValueError, "deve conter apenas strings"):
            validar_argumentos_obrigatorios(
                "criar_planilha",
                {"nome_arquivo": "gastos", "colunas": ["Nome", {"a": 1}]}
            )

    def test_validar_argumentos_obrigatorios_colunas_com_item_vazio_levanta_value_error(self):
        """Item string vazia/só espaços em 'colunas' deve ser rejeitado."""
        for colunas_invalidas in (["Nome", ""], ["Nome", "   "]):
            with self.assertRaisesRegex(ValueError, "deve conter apenas strings"):
                validar_argumentos_obrigatorios(
                    "criar_planilha",
                    {"nome_arquivo": "gastos", "colunas": colunas_invalidas}
                )

    def test_validar_argumentos_obrigatorios_colunas_valida_nao_levanta_excecao(self):
        """Lista de strings não-vazias em 'colunas' segue aceita (regressão)."""
        validar_argumentos_obrigatorios(
            "criar_planilha",
            {"nome_arquivo": "gastos", "colunas": ["Nome", "Idade"]}
        )

    def test_validar_argumentos_obrigatorios_colunas_lista_de_dicts_levanta_value_error(self):
        """Caso real do bug (modo de falha 2): lista de dicts em 'colunas' deve
        ser rejeitada na validação, antes de chegar ao pandas.DataFrame."""
        with self.assertRaisesRegex(ValueError, "deve conter apenas strings"):
            validar_argumentos_obrigatorios(
                "editar_planilha",
                {
                    "nome_arquivo": "produtos",
                    "colunas": [
                        {"Model": "QFY000013", "product": "produto A", "english description": "", "NCM": "4602"},
                        {"Model": "QFY000014", "product": "produto B", "english description": "", "NCM": "6302"},
                    ],
                }
            )


class TestGerarNomeUnico(unittest.TestCase):
    """Testes para função gerar_nome_unico."""
    
    def setUp(self):
        """Configura pasta temporária para testes."""
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_pasta = os.environ.get("PASTA_ARQUIVOS_GERADOS")
        os.environ["PASTA_ARQUIVOS_GERADOS"] = self.temp_dir.name
    
    def tearDown(self):
        """Limpa pasta temporária."""
        self.temp_dir.cleanup()
        if self.original_pasta:
            os.environ["PASTA_ARQUIVOS_GERADOS"] = self.original_pasta
        else:
            os.environ.pop("PASTA_ARQUIVOS_GERADOS", None)
    
    def test_conflito_gera_sufixo(self):
        """Testa que conflito de nome gera sufixo _1."""
        # Criar primeiro arquivo manualmente
        pasta = garantir_pasta_arquivos()
        primeiro_arquivo = os.path.join(pasta, "relatorio.xlsx")
        with open(primeiro_arquivo, "w") as f:
            f.write("teste")
        
        # Chamar gerar_nome_unico novamente
        nome_unico = gerar_nome_unico("relatorio", ".xlsx")
        
        self.assertEqual(nome_unico, "relatorio_1.xlsx")


class TestFluxoConfirmacao(unittest.TestCase):
    """Testes para fluxo de confirmação com ambiguidade."""
    
    def test_cancelamento_automatico_ambiguidade_repetida(self):
        """Testa cancelamento automático após duas respostas ambíguas."""
        sessao = ChatSession()
        
        # Definir ação pendente
        tool_call = {"name": "criar_planilha", "arguments": {"nome_arquivo": "teste"}}
        sessao.definir_acao_pendente(tool_call)
        
        # Primeira resposta ambígua
        resultado1 = interpretar_confirmacao("talvez")
        self.assertIsNone(resultado1)
        sessao.tentativas_confirmacao_ambigua += 1
        
        # Segunda resposta ambígua
        resultado2 = interpretar_confirmacao("hummm")
        self.assertIsNone(resultado2)
        sessao.tentativas_confirmacao_ambigua += 1
        
        # Verificar que deve cancelar automaticamente (contador >= 2)
        self.assertGreaterEqual(sessao.tentativas_confirmacao_ambigua, 2)
        
        # Simular cancelamento automático
        if sessao.tentativas_confirmacao_ambigua >= 2:
            sessao.limpar_acao_pendente()
        
        self.assertFalse(sessao.tem_acao_pendente())
        self.assertEqual(sessao.tentativas_confirmacao_ambigua, 0)


class TestRegressao(unittest.TestCase):
    """Testes de regressão para bugs conhecidos."""
    
    def test_simulacao_documento_sem_escape_literal(self):
        """Testa que string de simulação de criar_documento não contém \\\" literal."""
        argumentos = {
            "nome_arquivo": "relatorio",
            "titulo": "Relatório Mensal",
            "conteudo": "Conteúdo de teste"
        }
        
        resultado = simular_execucao_ferramenta("criar_documento", argumentos)
        
        # Verificar que não há escape literal de aspas duplas
        self.assertNotIn('\\"', resultado)
        # A string deve conter aspas normais, não escaped
        self.assertIn("Relatório Mensal", resultado)
    
    def test_timeout_de_streaming_nao_faz_retry(self):
        """Testa que timeout de geração é propagado sem nova tentativa."""
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask
        from backend.core.llama_client import LlamaTimeoutError

        class ClienteComTimeout:
            def __init__(self):
                self.model = "qwen3.5:4b"
                self.chamadas = 0

            def chat_com_tools_stream(self, **kwargs):
                self.chamadas += 1
                raise LlamaTimeoutError("timeout de teste")

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                self.chamadas += 1
                raise LlamaTimeoutError("timeout de teste")

        cliente = ClienteComTimeout()
        runner = MariaRunner(cliente=cliente)
        task = MariaTask(999, "Timeout", "Teste", "Olá")

        with self.assertRaises(LlamaTimeoutError):
            runner._enviar_com_retry(ChatSession(), task)

        self.assertEqual(cliente.chamadas, 1)

    def test_montar_mensagens_com_reforco_mescla_system_existente(self):


        historico = [
            {"role": "system", "content": "PROMPT LONGO ORIGINAL"},
            {"role": "user", "content": "mensagem anterior"},
            {"role": "assistant", "content": "resposta anterior"},
        ]
        mensagens = _montar_mensagens_com_reforco(historico, "nova mensagem")

        systems = [m for m in mensagens if m["role"] == "system"]
        self.assertEqual(len(systems), 1)
        self.assertIn("PROMPT LONGO ORIGINAL", systems[0]["content"])
        self.assertEqual(mensagens[-1], {"role": "user", "content": "nova mensagem"})
        self.assertEqual(historico[0]["content"], "PROMPT LONGO ORIGINAL")  # historico não mutado

    def test_montar_mensagens_com_reforco_sem_system_previo(self):
        mensagens = _montar_mensagens_com_reforco(None, "mensagem")

        systems = [m for m in mensagens if m["role"] == "system"]
        self.assertEqual(len(systems), 1)
        self.assertIn("Você é MARIA", systems[0]["content"])

class TestSessionStorage(unittest.TestCase):
    """Testes para persistência de sessões em disco."""

    def setUp(self):
        """Configura pasta temporária isolada para testes."""
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_pasta = os.environ.get("PASTA_SESSOES")
        os.environ["PASTA_SESSOES"] = self.temp_dir.name

    def tearDown(self):
        """Limpa pasta temporária e restaura variável de ambiente."""
        self.temp_dir.cleanup()
        if self.original_pasta:
            os.environ["PASTA_SESSOES"] = self.original_pasta
        else:
            os.environ.pop("PASTA_SESSOES", None)

    def test_salvar_e_carregar_sessao(self):
        """Testa que uma sessão salva pode ser recarregada com o mesmo histórico."""
        sessao = ChatSession()
        sessao.adicionar_mensagem("user", "Olá")
        sessao.adicionar_mensagem("assistant", "Oi! Como posso ajudar?")

        caminho = salvar_sessao(sessao.to_dict(), "sessao_20260101_120000.json")
        self.assertTrue(os.path.exists(caminho))

        dados_carregados = carregar_sessao(caminho)
        sessao_restaurada = ChatSession.from_dict(dados_carregados)

        self.assertEqual(sessao_restaurada.contar_mensagens(), 2)
        self.assertEqual(sessao_restaurada.get_ultima_mensagem_usuario(), "Olá")

    def test_listar_sessoes_ordena_mais_recentes_primeiro(self):
        """Testa que listar_sessoes_salvas ordena por nome (timestamp) decrescente."""
        sessao = ChatSession()
        sessao.adicionar_mensagem("user", "teste")

        salvar_sessao(sessao.to_dict(), "sessao_20260101_100000.json")
        salvar_sessao(sessao.to_dict(), "sessao_20260102_100000.json")

        sessoes = listar_sessoes_salvas()

        self.assertEqual(len(sessoes), 2)
        self.assertEqual(sessoes[0]["nome_arquivo"], "sessao_20260102_100000.json")
        self.assertEqual(sessoes[1]["nome_arquivo"], "sessao_20260101_100000.json")
        self.assertEqual(sessoes[0]["qtd_mensagens"], 1)

    def test_listar_sessoes_ignora_arquivo_corrompido(self):
        """Testa que um arquivo de sessão corrompido é ignorado, não levanta exceção."""
        pasta = os.environ["PASTA_SESSOES"]
        os.makedirs(pasta, exist_ok=True)
        caminho_corrompido = os.path.join(pasta, "sessao_20260103_100000.json")
        with open(caminho_corrompido, "w", encoding="utf-8") as arquivo:
            arquivo.write("{ json inválido")

        sessoes = listar_sessoes_salvas()

        self.assertEqual(sessoes, [])

    def test_carregar_sessao_inexistente_levanta_value_error(self):
        """Testa que carregar uma sessão inexistente levanta ValueError."""
        with self.assertRaises(ValueError):
            carregar_sessao(os.path.join(os.environ["PASTA_SESSOES"], "nao_existe.json"))


class TestAcessoLeitura(unittest.TestCase):
    """Testes para as ferramentas de leitura (listar_arquivos, ler_documento)."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_pastas = os.environ.get("PASTAS_PERMITIDAS")
        os.environ["PASTAS_PERMITIDAS"] = self.temp_dir.name

    def tearDown(self):
        self.temp_dir.cleanup()
        if self.original_pastas:
            os.environ["PASTAS_PERMITIDAS"] = self.original_pastas
        else:
            os.environ.pop("PASTAS_PERMITIDAS", None)

    def test_traversal_negado(self):
        """Testa que um caminho fora das pastas permitidas é rejeitado."""
        with self.assertRaises(ValueError):
            resolver_caminho_permitido("../../etc/passwd")

    def test_listar_pasta_temporaria(self):
        """Testa listagem de arquivos em uma pasta permitida."""
        caminho_arquivo = os.path.join(self.temp_dir.name, "nota.txt")
        with open(caminho_arquivo, "w", encoding="utf-8") as f:
            f.write("conteúdo de teste")

        itens = listar_arquivos()

        nomes = [item["nome"] for item in itens]
        self.assertIn("nota.txt", nomes)

    def test_listar_pasta_vazia(self):
        """Testa que uma pasta permitida vazia retorna lista vazia."""
        self.assertEqual(listar_arquivos(), [])

    def test_listar_pasta_padrao_ausente_nao_gera_caminho_duplicado(self):
        """Testa que listar_arquivos() sem pasta explícita não duplica o
        caminho e cria a pasta padrão automaticamente se ainda não existir."""
        pasta_inexistente = os.path.join(self.temp_dir.name, "ainda_nao_criada")
        os.environ["PASTAS_PERMITIDAS"] = pasta_inexistente

        self.assertFalse(os.path.isdir(pasta_inexistente))

        itens = listar_arquivos()

        self.assertEqual(itens, [])
        self.assertTrue(os.path.isdir(pasta_inexistente))

    def test_leitura_truncada(self):
        """Testa que um arquivo maior que max_chars é truncado corretamente."""
        caminho_arquivo = os.path.join(self.temp_dir.name, "grande.txt")
        with open(caminho_arquivo, "w", encoding="utf-8") as f:
            f.write("a" * 100)

        doc = ler_documento("grande.txt", max_chars=10)

        self.assertTrue(doc["truncado"])
        self.assertEqual(len(doc["texto"]), 10)
        self.assertEqual(doc["total_chars"], 100)

    def test_leitura_nao_truncada(self):
        """Testa que um arquivo menor que max_chars não é marcado como truncado."""
        caminho_arquivo = os.path.join(self.temp_dir.name, "pequeno.txt")
        with open(caminho_arquivo, "w", encoding="utf-8") as f:
            f.write("texto curto")

        doc = ler_documento("pequeno.txt")

        self.assertFalse(doc["truncado"])
        self.assertEqual(doc["texto"], "texto curto")

    def test_extensao_nao_suportada(self):
        """Testa que uma extensão fora da lista permitida levanta ValueError."""
        caminho_arquivo = os.path.join(self.temp_dir.name, "programa.exe")
        with open(caminho_arquivo, "wb") as f:
            f.write(b"binario")

        with self.assertRaises(ValueError):
            ler_documento("programa.exe")

    def test_arquivo_inexistente_levanta_value_error(self):
        """Testa que ler um arquivo inexistente levanta ValueError."""
        with self.assertRaises(ValueError):
            ler_documento("nao_existe.txt")


class TestFerramentasLeitura(unittest.TestCase):
    """Testes para o executor de ferramentas de leitura (tools_schema)."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_pastas = os.environ.get("PASTAS_PERMITIDAS")
        os.environ["PASTAS_PERMITIDAS"] = self.temp_dir.name

    def tearDown(self):
        self.temp_dir.cleanup()
        if self.original_pastas:
            os.environ["PASTAS_PERMITIDAS"] = self.original_pastas
        else:
            os.environ.pop("PASTAS_PERMITIDAS", None)

    def test_executar_listar_arquivos_pasta_vazia(self):
        resultado = executar_ferramenta_leitura("listar_arquivos", {})
        self.assertIn("vazia", resultado.lower())

    def test_executar_resumir_documento(self):
        caminho_arquivo = os.path.join(self.temp_dir.name, "notas.txt")
        with open(caminho_arquivo, "w", encoding="utf-8") as f:
            f.write("Reunião marcada para sexta-feira.")

        resultado = executar_ferramenta_leitura(
            "resumir_documento", {"nome_arquivo": "notas.txt"}
        )

        self.assertIn("Reunião marcada para sexta-feira.", resultado)

    def test_executar_ferramenta_desconhecida(self):
        with self.assertRaises(ValueError):
            executar_ferramenta_leitura("apagar_tudo", {})


class TestSelecaoModeloCLI(unittest.TestCase):
    """Testes para a seleção de modelo via argumento -m/--modelo (main.py)."""

    def test_controller_usa_modelo_informado(self):
        from backend.main import MariaController
        controller = MariaController(modelo="qwen3:8b")
        controller.inicializar()
        self.assertEqual(controller.cliente.model, "qwen3:8b")

    def test_controller_usa_modelo_padrao_quando_nao_informado(self):
        from backend.main import MariaController
        from backend.core.config import LLAMA_MODEL
        controller = MariaController()
        controller.inicializar()
        self.assertEqual(controller.cliente.model, LLAMA_MODEL)


class TestAquecimentoModelo(unittest.TestCase):
    """Testes para o warmup do modelo na inicialização interativa (main.py)."""

    def test_aquecer_modelo_chama_enviar_mensagem(self):
        from backend.main import MariaController
        controller = MariaController()
        controller.inicializar()
        controller.cliente.enviar_mensagem = MagicMock(return_value="ok")

        controller.aquecer_modelo()

        controller.cliente.enviar_mensagem.assert_called_once()

    def test_aquecer_modelo_nao_propaga_excecao(self):
        from backend.main import MariaController
        from backend.core.llama_client import LlamaClientError
        controller = MariaController()
        controller.inicializar()
        controller.cliente.enviar_mensagem = MagicMock(side_effect=LlamaClientError("falha"))

        controller.aquecer_modelo()  # não deve levantar exceção


class TestAcuraciaDeArgumentos(unittest.TestCase):
    """Testes para a comparação de argumentos do benchmark (MariaRunner)."""

    def test_argumentos_compativeis_sem_criterio_retorna_true(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        self.assertTrue(MariaRunner._argumentos_compativeis({"a": 1}, None))

    def test_argumentos_compativeis_subconjunto_correto(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        obtidos = {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"], "descricao": "extra"}
        esperados = {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]}
        self.assertTrue(MariaRunner._argumentos_compativeis(obtidos, esperados))

    def test_argumentos_incompativeis_detecta_divergencia(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        obtidos = {"nome_arquivo": "gastos_errado", "colunas": ["Data", "Valor"]}
        esperados = {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]}
        self.assertFalse(MariaRunner._argumentos_compativeis(obtidos, esperados))


class TestSystemPromptExcecaoArquivoInexistente(unittest.TestCase):
    """Testa que o SYSTEM_PROMPT distingue arquivo incerto de arquivo
    declaradamente inexistente (Fix D)."""

    def test_system_prompt_contem_excecao_para_arquivo_ficticio(self):
        self.assertIn("arquivo não foi encontrado", ChatSession.SYSTEM_PROMPT)


class TestToolChaining(unittest.TestCase):
    """Testes para o módulo compartilhado core/tool_chaining.py."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_pastas = os.environ.get("PASTAS_PERMITIDAS")
        os.environ["PASTAS_PERMITIDAS"] = self.temp_dir.name

    def tearDown(self):
        self.temp_dir.cleanup()
        if self.original_pastas:
            os.environ["PASTAS_PERMITIDAS"] = self.original_pastas
        else:
            os.environ.pop("PASTAS_PERMITIDAS", None)

    def test_encadear_leitura_stream_avanca_ate_ferramenta_de_escrita(self):
        from backend.core.tool_chaining import encadear_leitura_stream

        class ClienteFalso:
            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                yield None, {"name": "editar_planilha", "arguments": {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]}}

        resultado = list(encadear_leitura_stream(
            ClienteFalso(),
            historico_com_system=[{"role": "system", "content": "sistema"}],
            tool_call_inicial={"name": "listar_arquivos", "arguments": {}},
            tools=[],
        ))

        tool_call_final = resultado[-1][1]
        self.assertEqual(tool_call_final["name"], "editar_planilha")

    def test_encadear_leitura_stream_respeita_limite_de_passos(self):
        from backend.core.tool_chaining import encadear_leitura_stream

        class ClienteQueSempreLista:
            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                yield None, {"name": "listar_arquivos", "arguments": {}}

        resultado = list(encadear_leitura_stream(
            ClienteQueSempreLista(),
            historico_com_system=[{"role": "system", "content": "sistema"}],
            tool_call_inicial={"name": "listar_arquivos", "arguments": {}},
            tools=[],
        ))

        tool_call_final = resultado[-1][1]
        self.assertIsNone(tool_call_final)
        textos = "".join(chunk for chunk, _ in resultado if chunk)
        self.assertIn("Não consegui concluir", textos)

    def test_encadear_leitura_stream_propaga_timeout_por_chamada(self):
        from backend.core.tool_chaining import encadear_leitura_stream

        class ClienteFalso:
            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                yield None, {"name": "editar_planilha", "arguments": {}}

        def callback_que_estoura(duracao, tokens):
            raise TimeoutError("timeout de teste")

        with self.assertRaises(TimeoutError):
            list(encadear_leitura_stream(
                ClienteFalso(),
                historico_com_system=[{"role": "system", "content": "sistema"}],
                tool_call_inicial={"name": "listar_arquivos", "arguments": {}},
                tools=[],
                apos_cada_chamada=callback_que_estoura,
            ))

    def test_validar_e_corrigir_leitura_passa_direto_sem_chamar_continuacao(self):
        from backend.core.tool_chaining import validar_e_corrigir_tool_call_stream

        class ClienteFalso:
            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                raise AssertionError("não deveria ser chamado para leitura")

        resultado = list(validar_e_corrigir_tool_call_stream(
            ClienteFalso(),
            historico_com_system=[{"role": "system", "content": "sistema"}],
            tool_call_atual={"name": "listar_arquivos", "arguments": {}},
            tools=[],
        ))
        ultimo = resultado[-1][1]
        self.assertEqual(ultimo["tool_call"]["name"], "listar_arquivos")
        self.assertEqual(ultimo["tentativas"], 0)

    def test_validar_e_corrigir_escrita_invalida_corrigida_na_primeira_tentativa(self):
        from backend.core.tool_chaining import validar_e_corrigir_tool_call_stream

        class ClienteFalso:
            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                yield None, {
                    "name": "criar_planilha",
                    "arguments": {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]},
                }

        resultado = list(validar_e_corrigir_tool_call_stream(
            ClienteFalso(),
            historico_com_system=[{"role": "system", "content": "sistema"}],
            tool_call_atual={
                "name": "criar_planilha",
                "arguments": {"nome_arquivo": "gastos", "colunas": "Data,Valor"},
            },
            tools=[],
        ))
        ultimo = resultado[-1][1]
        self.assertEqual(ultimo["tool_call"]["name"], "criar_planilha")
        self.assertEqual(ultimo["tool_call"]["arguments"]["colunas"], ["Data", "Valor"])
        self.assertEqual(ultimo["tentativas"], 1)

    def test_validar_e_corrigir_escrita_esgota_limite_sem_correcao(self):
        from backend.core.tool_chaining import validar_e_corrigir_tool_call_stream

        class ClienteSempreInvalido:
            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                yield None, {
                    "name": "criar_planilha",
                    "arguments": {"nome_arquivo": "gastos", "colunas": "Data,Valor"},
                }

        resultado = list(validar_e_corrigir_tool_call_stream(
            ClienteSempreInvalido(),
            historico_com_system=[{"role": "system", "content": "sistema"}],
            tool_call_atual={
                "name": "criar_planilha",
                "arguments": {"nome_arquivo": "gastos", "colunas": "Data,Valor"},
            },
            tools=[],
        ))
        ultimo = resultado[-1][1]
        self.assertIsNone(ultimo["tool_call"])
        self.assertEqual(ultimo["tentativas"], 2)


class TestMariaRunnerEncadeamento(unittest.TestCase):
    """Testa que o MariaRunner encadeia leitura -> escrita (Fix A)."""

    def test_runner_encadeia_listar_arquivos_ate_editar_planilha(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteFalso:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return ("", {"name": "listar_arquivos", "arguments": {}}, 10, 5.0, 1.0)

            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                yield None, {
                    "name": "editar_planilha",
                    "arguments": {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]},
                }

        task = MariaTask(
            9001, "Teste encadeamento", "desc", "edite a planilha gastos",
            expected_tool="editar_planilha", confirm_sequence=[],
            category=MariaTaskCategory.EDITAR_PLANILHA,
        )

        runner = MariaRunner(cliente=ClienteFalso())
        resultado = runner.run(task)

        self.assertEqual(resultado.tool_detected, "editar_planilha")
        self.assertTrue(resultado.tool_correct)

    def test_runner_corrige_tool_call_escrita_invalida(self):
        """Tarefa com tool call de escrita inválida (schema) deve registrar correction_attempts > 0."""
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteCorrige:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                # Primeira chamada: criar_planilha com 'colunas' como string (inválida).
                return (
                    "", 
                    {"name": "criar_planilha", "arguments": {"nome_arquivo": "gastos", "colunas": "Data,Valor"}},
                    10, 5.0, 1.0,
                )

            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                # Correção: produz uma tool call válida.
                yield None, {
                    "name": "criar_planilha",
                    "arguments": {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]},
                }

        task = MariaTask(
            9002, "Teste correção tool call", "desc", "crie a planilha gastos",
            expected_tool="criar_planilha", confirm_sequence=[],
            category=MariaTaskCategory.CRIAR_PLANILHA,
        )

        runner = MariaRunner(cliente=ClienteCorrige())
        resultado = runner.run(task)

        self.assertEqual(resultado.tool_detected, "criar_planilha")
        self.assertTrue(resultado.tool_correct)
        self.assertGreater(resultado.correction_attempts, 0)


class TestOrcamentoDeTokensParaDocumento(unittest.TestCase):
    """Testa a heurística de orçamento maior de tokens para composição de
    documentos narrativos (Item A)."""

    def test_sugere_composicao_de_documento_detecta_palavras_chave(self):
        from backend.core.llama_client import _sugere_composicao_de_documento
        self.assertTrue(_sugere_composicao_de_documento("Crie uma carta de apresentação formal"))
        self.assertTrue(_sugere_composicao_de_documento("Escreva um relatório da reunião"))
        self.assertFalse(_sugere_composicao_de_documento("Crie uma planilha de gastos"))

class TestEncadeamentoPropagaTokens(unittest.TestCase):
    """Testa que core/tool_chaining.py repassa tokens_gerados da chamada de
    continuação para o callback apos_cada_chamada (Item B)."""

    def test_apos_cada_chamada_recebe_tokens_gerados(self):
        from backend.core.tool_chaining import encadear_leitura_stream

        class ClienteFalso:
            def continuar_com_resultado_ferramenta_stream(self, metricas_saida=None, **kwargs):
                if metricas_saida is not None:
                    metricas_saida["tokens_gerados"] = 42
                yield None, {"name": "editar_planilha", "arguments": {}}

        chamadas = []
        list(encadear_leitura_stream(
            ClienteFalso(),
            historico_com_system=[{"role": "system", "content": "sistema"}],
            tool_call_inicial={"name": "listar_arquivos", "arguments": {}},
            tools=[],
            apos_cada_chamada=lambda duracao, tokens: chamadas.append(tokens),
        ))

        self.assertEqual(chamadas, [42])


class TestMariaRunnerSomaTokensDaContinuacao(unittest.TestCase):
    """Testa que MariaRunner soma tokens_gerados da chamada de continuação
    ao total reportado (Item B)."""

    def test_tokens_gerados_inclui_continuacao(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteFalso:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return ("", {"name": "listar_arquivos", "arguments": {}}, 10, 5.0, 1.0)

            def continuar_com_resultado_ferramenta_stream(self, metricas_saida=None, **kwargs):
                if metricas_saida is not None:
                    metricas_saida["tokens_gerados"] = 15
                yield None, {
                    "name": "editar_planilha",
                    "arguments": {"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]},
                }

        task = MariaTask(
            9002, "Teste soma tokens", "desc", "edite a planilha gastos",
            expected_tool="editar_planilha", confirm_sequence=[],
            category=MariaTaskCategory.EDITAR_PLANILHA,
        )

        runner = MariaRunner(cliente=ClienteFalso())
        resultado = runner.run(task)

        self.assertEqual(resultado.tokens_gerados, 25)  # 10 (inicial) + 15 (continuação)


class TestRunRepeatedComCallback(unittest.TestCase):
    """Testa que MariaRunner.run_repeated invoca o callback de progresso
    após cada execução individual (Item C)."""

    def test_apos_cada_execucao_e_chamado_para_cada_repeticao(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteFalso:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return ("Olá!", None, 5, 10.0, 0.5)

        task = MariaTask(9003, "Teste callback", "desc", "olá", category=MariaTaskCategory.CONVERSA)
        runner = MariaRunner(cliente=ClienteFalso())

        chamadas = []
        runner.run_repeated(task, 3, apos_cada_execucao=lambda i, r: chamadas.append(i))

        self.assertEqual(chamadas, [1, 2, 3])


class TestMariaRunnerNegaEAmbiguidade(unittest.TestCase):
    """Tarefa 4: em negação ou ambiguidade, a ferramenta não é executada e
    tool_call_final passa a ser None (tool_correct=True quando
    expected_tool=None)."""

    def _falso_cliente_com_tool(self, nome_tool: str):
        class ClienteFalso:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return ("", {"name": nome_tool, "arguments": {}}, 5, 2.0, 0.5)

            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                # Correção da tool call inválida: devolve chamada válida na
                # primeira tentativa (integração correção <-> cancelamento).
                argumentos_validos = {
                    "criar_planilha": {"nome_arquivo": "teste", "colunas": ["Data"]},
                    "criar_documento": {"nome_arquivo": "teste", "titulo": "T", "conteudo": "C"},
                    "editar_planilha": {"nome_arquivo": "teste", "colunas": ["Data"]},
                }
                yield None, {"name": nome_tool, "arguments": argumentos_validos[nome_tool]}

        return ClienteFalso()

    def test_negacao_anula_tool_call(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        task = MariaTask(
            9004, "Negação teste", "desc",
            "Crie uma planilha simples de tarefas.",
            expected_tool=None, expected_keywords=["cancelada"],
            confirm_sequence=["não"], expected_final_message="cancelada",
            category=MariaTaskCategory.CANCELAMENTO,
        )
        runner = MariaRunner(cliente=self._falso_cliente_com_tool("criar_planilha"))
        resultado = runner.run(task)

        self.assertIsNone(resultado.tool_detected)
        self.assertTrue(resultado.tool_correct)
        self.assertTrue(resultado.confirmation_completed)
        self.assertEqual(resultado.final_message, "Ação cancelada.")

    def test_ambiguidade_anula_tool_call(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        task = MariaTask(
            9005, "Ambiguidade teste", "desc",
            "Crie uma planilha de projetos com Projeto e Status.",
            expected_tool=None, expected_keywords=["cancelada"],
            confirm_sequence=["talvez", "hummm"], expected_final_message="cancelada",
            category=MariaTaskCategory.AMBIGUIDADE,
        )
        runner = MariaRunner(cliente=self._falso_cliente_com_tool("criar_planilha"))
        resultado = runner.run(task)

        self.assertIsNone(resultado.tool_detected)
        self.assertTrue(resultado.tool_correct)
        self.assertTrue(resultado.confirmation_completed)
        self.assertEqual(resultado.final_message, "Ação cancelada por ambiguidade.")


class TestMariaRunnerMensagemDeErro(unittest.TestCase):
    """Arquivo inexistente: após a confirmação, a ferramenta EXECUTA e falha
    (ValueError); o erro REAL é devolvido ao modelo via continuação e a
    resposta do modelo vira a mensagem final — não é erro de tarefa."""

    def test_value_error_edicao_inexistente_devolvido_ao_modelo(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        continuacao_vista = {}

        class ClienteFalso:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return ("", {
                    "name": "editar_planilha",
                    "arguments": {"nome_arquivo": "arquivo_ausente_a", "colunas": ["A"]},
                }, 5, 2.0, 0.5)

            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                continuacao_vista.update(kwargs)
                yield "O arquivo não existe. Deseja criar uma nova?", None

        task = MariaTask(
            9006, "Edição inexistente teste", "desc",
            "Edite a planilha arquivo_ausente_a com a coluna A.",
            expected_tool=None, expected_keywords=["exist"],
            confirm_sequence=["sim"], category=MariaTaskCategory.EDITAR_PLANILHA,
        )
        runner = MariaRunner(cliente=ClienteFalso())
        resultado = runner.run(task)

        # O erro real da ferramenta voltou ao modelo (não virou erro de tarefa).
        self.assertTrue(resultado.runtime_ok)
        self.assertFalse(resultado.errors)
        self.assertIn("não encontrado", continuacao_vista["resultado"])
        self.assertEqual(continuacao_vista["tool_call"]["name"], "editar_planilha")
        self.assertIn("não existe", resultado.final_message)



class TestMariaRunnerCadeiaFerramentas(unittest.TestCase):
    """Tarefas com tools_obrigatorios=["editar_planilha"]: a ferramenta precisa
    ter sido chamada na cadeia E a execução precisa terminar em texto — simula
    2 turnos reais: usuário pede edição → modelo chama editar_planilha → a
    ferramenta EXECUTA e falha (arquivo ausente) → o erro volta ao modelo →
    ele responde em texto."""

    def setUp(self):
        """Garante a pasta de arquivos do benchmark vazia: as tarefas sob teste
        exigem que o arquivo NÃO exista para a ferramenta falhar em runtime."""
        import shutil
        from backend.benchmarks.maria_bench.benchmark_config import BENCHMARK_ARQUIVOS_DIR
        if os.path.isdir(BENCHMARK_ARQUIVOS_DIR):
            for item in os.listdir(BENCHMARK_ARQUIVOS_DIR):
                item_path = os.path.join(BENCHMARK_ARQUIVOS_DIR, item)
                try:
                    if os.path.isfile(item_path):
                        os.unlink(item_path)
                    elif os.path.isdir(item_path):
                        shutil.rmtree(item_path)
                except OSError:
                    pass

    def _task(self, **kwargs):
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        base = dict(
            id=9010,
            name="Edição inexistente teste",
            description="desc",
            user_message="Edite a planilha estoque com a coluna preco.",
            tools_obrigatorios=["editar_planilha"],
            expected_keywords=["nao existe", "nao encontrado", "encontrado", "vazia", "criar"],
            confirm_sequence=["sim"],
            category=MariaTaskCategory.EDITAR_PLANILHA,
        )
        base.update(kwargs)
        return MariaTask(**base)

    def test_erro_da_ferramenta_devolvido_ao_modelo_que_responde_em_texto_conta_como_correta(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner

        continuacao_vista = {}

        class ClienteEditaInexistente:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return (
                    "",
                    {"name": "editar_planilha", "arguments": {"nome_arquivo": "estoque", "colunas": ["preco"]}},
                    10, 5.0, 1.0,
                )

            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                continuacao_vista.update(kwargs)
                yield "A planilha estoque não existe. Deseja criar uma nova?", None

        resultado = MariaRunner(cliente=ClienteEditaInexistente()).run(self._task())

        self.assertIsNone(resultado.tool_detected)
        self.assertTrue(resultado.tool_correct)
        self.assertEqual(resultado.cadeia_ferramentas, ["editar_planilha"])
        self.assertEqual(resultado.tool_call_inicial.get("name"), "editar_planilha")
        self.assertTrue(resultado.confirmation_completed)
        self.assertTrue(resultado.runtime_ok)
        self.assertTrue(resultado.keyword_match)
        # O erro REAL da ferramenta (arquivo não encontrado) voltou ao modelo.
        self.assertIn("não encontrado", continuacao_vista["resultado"])
        self.assertEqual(continuacao_vista["tool_call"]["name"], "editar_planilha")
        self.assertIn("não existe", resultado.final_message)

    def test_rechamada_da_ferramenta_apos_erro_conta_como_incorreta(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner

        class ClienteTeimoso:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return (
                    "",
                    {"name": "editar_planilha", "arguments": {"nome_arquivo": "estoque", "colunas": ["preco"]}},
                    10, 5.0, 1.0,
                )

            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                yield None, {
                    "name": "editar_planilha",
                    "arguments": {"nome_arquivo": "estoque", "colunas": ["preco"]},
                }

        resultado = MariaRunner(cliente=ClienteTeimoso()).run(self._task())

        self.assertFalse(resultado.tool_correct)
        self.assertEqual(resultado.tool_detected, "editar_planilha")
        self.assertEqual(resultado.cadeia_ferramentas, ["editar_planilha"])

    def test_resposta_em_texto_sem_chamar_ferramenta_conta_como_incorreta(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner

        class ClienteNaoChama:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return ("Não encontrei a planilha em nenhum lugar.", None, 10, 5.0, 1.0)

        resultado = MariaRunner(cliente=ClienteNaoChama()).run(self._task())

        self.assertFalse(resultado.tool_correct)
        self.assertEqual(resultado.cadeia_ferramentas, [])

    def test_ferramenta_leitura_intermediaria_entra_na_cadeia(self):
        """FIX-4: ferramenta de leitura chamada DENTRO do encadear_leitura_stream
        (passo intermediário) entra em cadeia_ferramentas — não apenas a inicial
        (linha ~184) e a final (linha ~381). Fluxo simulado: listar_arquivos
        (inicial) → extrair_dados_planilha (intermediária) → criar_planilha (final)."""
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteLeituraEscrita:
            model = "modelo-teste"

            def __init__(self):
                self._chamadas_continuacao = 0

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                # Ferramenta INICIAL de leitura: dispara o encadeamento.
                return (
                    "",
                    {"name": "listar_arquivos", "arguments": {"pasta": "arquivos_gerados"}},
                    10, 5.0, 1.0,
                )

            def continuar_com_resultado_ferramenta_stream(self, **kwargs):
                # 1ª continuação → nova leitura (extrair_dados_planilha): passo
                # INTERMEDIÁRIO do encadeamento.
                # 2ª continuação → ferramenta de escrita (criar_planilha): fim.
                self._chamadas_continuacao += 1
                if self._chamadas_continuacao == 1:
                    yield None, {
                        "name": "extrair_dados_planilha",
                        "arguments": {"nome_arquivo": "produtos_mandarim"},
                    }
                else:
                    yield None, {
                        "name": "criar_planilha",
                        "arguments": {"nome_arquivo": "resultado", "colunas": ["A"]},
                    }

        task = MariaTask(
            id=9020,
            name="Leitura encadeada com passo intermediário",
            description="desc",
            user_message=(
                "Liste os arquivos, leia produtos_mandarim e crie a planilha resultado."
            ),
            tools_obrigatorios=["extrair_dados_planilha", "criar_planilha"],
            expected_tool="criar_planilha",
            confirm_sequence=["sim"],
            category=MariaTaskCategory.CRIAR_PLANILHA,
            expected_args_subset={"nome_arquivo": "resultado"},
        )

        with patch(
            "backend.benchmarks.maria_bench.runners.maria_runner.executar_ferramenta_real",
            return_value="Planilha criada com sucesso: resultado.xlsx",
        ), patch(
            "backend.application.tool_chaining.executar_ferramenta_leitura",
            return_value="arquivo1.xlsx",
        ):
            resultado = MariaRunner(cliente=ClienteLeituraEscrita()).run(task)

        # extrair_dados_planilha (intermediária) só entra na cadeia POR CAUSA do
        # callback apos_cada_leitura — sem o FIX-4, a cadeia seria inicial+final
        # (["listar_arquivos", "criar_planilha"]) e tool_correct seria False.
        self.assertEqual(
            resultado.cadeia_ferramentas,
            ["listar_arquivos", "extrair_dados_planilha", "criar_planilha"],
        )
        self.assertTrue(resultado.tool_correct)


class TestEncadearLeituraStreamCallback(unittest.TestCase):
    """FIX-4: o callback apos_cada_leitura é chamado com nome e argumentos de
    cada ferramenta de leitura executada no encadeamento."""

    def test_callback_chamado_para_cada_leitura(self):
        from backend.application.tool_chaining import encadear_leitura_stream

        chamadas = []

        def callback_leitura(nome: str, argumentos: dict) -> None:
            chamadas.append((nome, argumentos))

        cliente = MagicMock()
        # 1ª continuação: retorna outra leitura (listar_arquivos) → nova iteração.
        # 2ª continuação: texto → encerra o encadeamento.
        cliente.continuar_com_resultado_ferramenta_stream.side_effect = [
            iter([
                (None, {"name": "listar_arquivos", "arguments": {"pasta": "x"}}),
            ]),
            iter([("Resposta final.", None)]),
        ]

        tool_inicial = {"name": "listar_arquivos", "arguments": {"pasta": "arquivos_gerados"}}

        with patch(
            "backend.application.tool_chaining.executar_ferramenta_leitura",
            return_value="arquivo1.xlsx",
        ):
            list(encadear_leitura_stream(
                cliente=cliente,
                historico_com_system=[],
                tool_call_inicial=tool_inicial,
                tools=[],
                apos_cada_leitura=callback_leitura,
            ))

        self.assertTrue(chamadas, "callback nunca foi chamado")
        nomes_chamados = [nome for nome, _ in chamadas]
        self.assertIn("listar_arquivos", nomes_chamados)
        # Argumentos repassados ao callback (da tool call inicial).
        self.assertIn({"pasta": "arquivos_gerados"}, [args for _, args in chamadas])

    def test_sem_callback_nao_levanta_excecao(self):
        from backend.application.tool_chaining import encadear_leitura_stream

        cliente = MagicMock()
        cliente.continuar_com_resultado_ferramenta_stream.return_value = iter([
            ("Resposta.", None),
        ])

        tool_inicial = {"name": "listar_arquivos", "arguments": {}}

        with patch(
            "backend.application.tool_chaining.executar_ferramenta_leitura",
            return_value="ok",
        ):
            resultado = list(encadear_leitura_stream(
                cliente=cliente,
                historico_com_system=[],
                tool_call_inicial=tool_inicial,
                tools=[],
                apos_cada_leitura=None,  # sem callback
            ))

        self.assertIsNotNone(resultado)


class TestTarefas22E23EscritaInexistente(unittest.TestCase):
    """Tarefas 22 e 23: mensagem realista (não entrega a inexistência ao
    modelo); ele deve CHAMAR editar_planilha (que executa e falha com arquivo
    ausente, devolvendo o erro ao modelo) e responder em texto."""

    def test_desenho_das_tarefas_22_e_23(self):
        from backend.benchmarks.maria_bench.tasks import load_all_maria_tasks

        tarefas = {t.id: t for t in load_all_maria_tasks()}
        for tid in (22, 23):
            task = tarefas[tid]
            self.assertEqual(task.tools_obrigatorios, ["editar_planilha"])
            self.assertEqual(task.confirm_sequence, ["sim"])
            self.assertEqual(task.fixtures, [])
            self.assertIsNone(task.expected_tool)
            self.assertTrue(task.expected_keywords)
            mensagem = task.user_message.lower()
            self.assertNotIn("inexistente", mensagem)
            self.assertNotIn("nao existe", mensagem)

# ═══════════════════════════════════════════════════════════════
# Testes do LlamaClient (llama-server / API OpenAI-compatible)
# ═══════════════════════════════════════════════════════════════

class TestLlamaClientChatTexto(unittest.TestCase):
    """Testa chat síncrono básico do LlamaClient."""

    @patch('backend.core.llama_client.requests.Session')
    def test_chat_retorna_texto(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)
        mock_session.post.return_value = MagicMock(
            status_code=200,
            json=lambda: {
                "choices": [{"message": {"content": "Olá!", "tool_calls": None}}],
                "usage": {"completion_tokens": 5},
            },
        )

        cliente = LlamaClient()
        texto, tool_call = cliente.chat([{"role": "user", "content": "Oi"}])

        self.assertEqual(texto, "Olá!")
        self.assertIsNone(tool_call)

    @patch('backend.core.llama_client.requests.Session')
    def test_chat_preenche_metricas(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)
        mock_session.post.return_value = MagicMock(
            status_code=200,
            json=lambda: {
                "choices": [{"message": {"content": "ok", "tool_calls": None}}],
                "usage": {"completion_tokens": 7},
            },
        )

        cliente = LlamaClient()
        metricas = {}
        cliente.chat([{"role": "user", "content": "teste"}], metricas_saida=metricas)

        self.assertEqual(metricas["tokens_gerados"], 7)


class TestLlamaClientToolCalling(unittest.TestCase):
    """Testa tool calling estruturado e fallback textual."""

    @patch('backend.core.llama_client.requests.Session')
    def test_tool_call_estruturada(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)
        mock_session.post.return_value = MagicMock(
            status_code=200,
            json=lambda: {
                "choices": [{
                    "message": {
                        "content": "",
                        "tool_calls": [{
                            "function": {
                                "name": "criar_planilha",
                                "arguments": '{"nome_arquivo": "gastos", "colunas": ["Data", "Valor"]}',
                            }
                        }],
                    }
                }],
                "usage": {"completion_tokens": 20},
            },
        )

        cliente = LlamaClient()
        _, tool_call = cliente.chat([{"role": "user", "content": "crie planilha"}], tools=[{}])

        self.assertIsNotNone(tool_call)
        self.assertEqual(tool_call["name"], "criar_planilha")
        self.assertEqual(tool_call["arguments"]["nome_arquivo"], "gastos")

    @patch('backend.core.llama_client.requests.Session')
    def test_fallback_textual_extrai_tool_call(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        conteudo_com_tool_call = '{"name": "criar_planilha", "arguments": {"nome_arquivo": "test"}}'
        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)
        mock_session.post.return_value = MagicMock(
            status_code=200,
            json=lambda: {
                "choices": [{"message": {"content": conteudo_com_tool_call, "tool_calls": None}}],
                "usage": {"completion_tokens": 15},
            },
        )

        cliente = LlamaClient()
        _, tool_call = cliente.chat([{"role": "user", "content": "crie planilha"}], tools=[{}])

        self.assertIsNotNone(tool_call)
        self.assertEqual(tool_call["name"], "criar_planilha")


class TestLlamaClientStreaming(unittest.TestCase):
    """Testa streaming do LlamaClient."""

    @patch('backend.core.llama_client.requests.Session')
    def test_chat_stream_yields_chunks(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)

        linhas = [
            b'data: ' + json.dumps({"choices": [{"delta": {"content": "Olá"}, "finish_reason": None}]}).encode(),
            b'data: ' + json.dumps({"choices": [{"delta": {"content": " mundo"}, "finish_reason": None}]}).encode(),
            b'data: ' + json.dumps({"choices": [{"delta": {}, "finish_reason": "stop"}], "usage": {"completion_tokens": 2}}).encode(),
            b'data: [DONE]',
        ]
        mock_response = MagicMock(status_code=200)
        mock_response.iter_lines.return_value = iter(linhas)
        mock_session.post.return_value = mock_response

        cliente = LlamaClient()
        chunks = list(cliente.chat_stream([{"role": "user", "content": "oi"}]))

        textos = [c for c, _ in chunks if c is not None]
        self.assertIn("Olá", textos)
        self.assertIn(" mundo", textos)
        # Último item deve ser (None, None) ou (None, tool_call)
        ultimo_chunk, ultimo_tool = chunks[-1]
        self.assertIsNone(ultimo_chunk)

    @patch('backend.core.llama_client.requests.Session')
    def test_chat_stream_metricas(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)

        linhas = [
            b'data: ' + json.dumps({"choices": [{"delta": {"content": "ok"}, "finish_reason": None}]}).encode(),
            b'data: ' + json.dumps({"choices": [{"delta": {}, "finish_reason": "stop"}], "usage": {"completion_tokens": 3}}).encode(),
        ]
        mock_response = MagicMock(status_code=200)
        mock_response.iter_lines.return_value = iter(linhas)
        mock_session.post.return_value = mock_response

        cliente = LlamaClient()
        metricas = {}
        list(cliente.chat_stream([{"role": "user", "content": "oi"}], metricas_saida=metricas))

        self.assertEqual(metricas["tokens_gerados"], 3)
        self.assertIn("ttft", metricas)

    @patch('backend.core.llama_client.requests.Session')
    def test_chat_com_tools_stream_com_metricas_retorna_metricas(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)

        linhas = [
            b'data: ' + json.dumps({"choices": [{"delta": {"content": "R"}, "finish_reason": None}]}).encode(),
            b'data: ' + json.dumps({"choices": [{"delta": {"content": "esposta"}, "finish_reason": None}]}).encode(),
            b'data: ' + json.dumps({"choices": [{"delta": {}, "finish_reason": "stop"}], "usage": {"completion_tokens": 2}}).encode(),
            b'data: [DONE]',
        ]
        mock_response = MagicMock(status_code=200)
        mock_response.iter_lines.return_value = iter(linhas)
        mock_session.post.return_value = mock_response

        cliente = LlamaClient()
        self.assertTrue(hasattr(cliente, "chat_com_tools_stream_com_metricas"))

        texto, tool_call, tokens_gerados, tokens_por_segundo, ttft_ms = (
            cliente.chat_com_tools_stream_com_metricas(
                mensagem_usuario="olá",
                historico=[],
                tools=None,
            )
        )
        self.assertEqual(texto, "Resposta")
        self.assertIsNone(tool_call)
        self.assertGreater(tokens_gerados, 0)
        self.assertGreater(tokens_por_segundo, 0)
        self.assertIsNotNone(ttft_ms)


class TestLlamaClientErros(unittest.TestCase):
    """Testa tratamento de erros de conexão e timeout."""

    @patch('backend.core.llama_client.requests.Session')
    def test_erro_conexao_levanta_llama_client_error(self, mock_session_class):
        import requests as req
        from backend.core.llama_client import LlamaClient, LlamaClientError

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.side_effect = req.exceptions.ConnectionError("offline")

        cliente = LlamaClient()
        with self.assertRaises(LlamaClientError):
            cliente.chat([{"role": "user", "content": "oi"}])

    @patch('backend.core.llama_client.requests.Session')
    def test_timeout_levanta_llama_timeout_error(self, mock_session_class):
        import requests as req
        from backend.core.llama_client import LlamaClient, LlamaTimeoutError

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)
        mock_session.post.side_effect = req.exceptions.Timeout("timeout")

        cliente = LlamaClient()
        with self.assertRaises(LlamaTimeoutError):
            cliente.chat([{"role": "user", "content": "oi"}])

class TestDeteccaoDegeneracao(unittest.TestCase):
    """Testa a detecção de loop degenerado de geração (ex.: '\\n' x 600)."""

    def test_repeticao_de_newlines_detectada(self):
        from backend.core.llama_client import _detectar_degeneracao
        self.assertTrue(_detectar_degeneracao("texto: " + "\n" * 100))

    def test_texto_normal_nao_detectado(self):
        from backend.core.llama_client import _detectar_degeneracao
        self.assertFalse(_detectar_degeneracao('criar_planilha: ["gastos", ["Data", "Valor"]]'))

    def test_texto_abaixo_do_limiar_nao_detectado(self):
        from backend.core.llama_client import _detectar_degeneracao
        self.assertFalse(_detectar_degeneracao("\n" * 99))

    def test_repeticao_de_bloco_dois_chars_detectada(self):
        # Padrão real da task 10 (run_20260905_101728): bloco "_x" repetido.
        from backend.core.llama_client import _detectar_degeneracao
        self.assertTrue(_detectar_degeneracao("texto: " + "_x" * 60))

    def test_repeticao_de_bloco_dois_chars_com_prefixo_valido(self):
        # Tool call textual válida seguida de loop "_x" (caso observado no log).
        from backend.core.llama_client import _detectar_degeneracao
        prefixo = 'criar_documento: ["Novo_horario", "Mudança_horario", "A partir_de'
        self.assertTrue(_detectar_degeneracao(prefixo + "_x" * 60))

    def test_repeticao_de_bloco_de_tres_chars_detectada(self):
        from backend.core.llama_client import _detectar_degeneracao
        self.assertTrue(_detectar_degeneracao("x: " + "abc" * 40))

    def test_bloco_dois_chars_abaixo_do_limiar_nao_detectado(self):
        from backend.core.llama_client import _detectar_degeneracao
        self.assertFalse(_detectar_degeneracao("_x" * 49))  # 98 chars < minimo

    def test_texto_real_longo_nao_detectado(self):
        # Conteúdo narrativo legítimo com 100+ chars não deve ser sinalizado.
        from backend.core.llama_client import _detectar_degeneracao
        texto = (
            "A diretoria informa que, a partir da próxima segunda-feira, o novo "
            "horário de funcionamento será das 8h às 17h, com intervalo para "
            "almoço das 12h às 13h, conforme deliberado em reunião."
        )
        self.assertFalse(_detectar_degeneracao(texto))


class TestAnaliseSemantica(unittest.TestCase):
    """Testa a análise semântica heurística da tool call final (Fase 2)."""

    def _analisar(self, tool_call):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        return MariaRunner._analisar_semantica(tool_call)

    def test_titulo_conteudo_invertido_detectado(self):
        flags = self._analisar({
            "name": "criar_documento",
            "arguments": {
                "nome_arquivo": "comunicado",
                "titulo": "Informamos que o horário de atendimento será alterado a partir da próxima semana para melhor atender nossos clientes.",
                "conteudo": "Agradecemos a compreensão.",
            },
        })
        self.assertTrue(flags["titulo_conteudo_invertido"])

    def test_placeholder_detectado(self):
        flags = self._analisar({
            "name": "criar_documento",
            "arguments": {
                "nome_arquivo": "ata",
                "titulo": "Ata",
                "conteudo": "Data: [data] Local: [local]",
            },
        })
        self.assertTrue(flags["placeholder_detectado"])

    def test_conteudo_curto_detectado(self):
        flags = self._analisar({
            "name": "criar_documento",
            "arguments": {"nome_arquivo": "a", "titulo": "T", "conteudo": "muito curto"},
        })
        self.assertTrue(flags["conteudo_curto"])

    def test_nome_com_extensao_detectado(self):
        flags = self._analisar({
            "name": "criar_planilha",
            "arguments": {"nome_arquivo": "estoque.xlsx", "colunas": ["A"]},
        })
        self.assertTrue(flags["nome_com_extensao"])

    def test_tool_call_limpa_sem_flags(self):
        flags = self._analisar({
            "name": "criar_documento",
            "arguments": {
                "nome_arquivo": "ata",
                "titulo": "Ata",
                "conteudo": "Reunião realizada com sucesso, com discussão dos pontos de pauta.",
            },
        })
        self.assertFalse(any(flags.values()))

    def test_sem_tool_call_retorna_flags_falsas(self):
        flags = self._analisar(None)
        self.assertFalse(any(flags.values()))

    def test_calculate_maria_metrics_com_semantica(self):
        from backend.benchmarks.maria_bench.analysis.metrics import calculate_maria_metrics
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTaskResult
        base = dict(task_name="t", category="c", model="m", tool_detected=None,
                    tool_correct=True, confirmation_completed=True,
                    keyword_match=True, runtime_ok=True, final_message="",
                    latency_ms=0.0)
        r_limpo = MariaTaskResult(task_id=1, **base)
        r_placeholder = MariaTaskResult(task_id=2, **base, placeholder_detectado=True)
        r_corrigido = MariaTaskResult(
            task_id=3, **base,
            correcoes=[{"campo": "nome_arquivo", "antes": "../x", "depois": "x"}],
        )
        metrics = calculate_maria_metrics([r_limpo, r_placeholder, r_corrigido])
        self.assertAlmostEqual(metrics.semantic_quality_rate, 2 / 3)
        self.assertEqual(metrics.semantic_errors_by_type.get("placeholder_detectado"), 1)
        self.assertEqual(metrics.correcoes_count, 1)


class TestFormatarAvisos(unittest.TestCase):
    """Testa a lista de avisos (⚠️ ...) exibida em linhas separadas."""

    def _resultado(self, **kwargs):
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTaskResult
        base = dict(task_id=1, task_name="t", category="c", model="m",
                    tool_detected=None, tool_correct=True,
                    confirmation_completed=True, keyword_match=True,
                    runtime_ok=True, final_message="", latency_ms=0.0)
        base.update(kwargs)
        return MariaTaskResult(**base)

    def test_sem_avisos_retorna_vazio(self):
        from backend.benchmarks.maria_bench.analysis.report import formatar_avisos
        self.assertEqual(formatar_avisos(self._resultado()), [])

    def test_correcao_gera_aviso_antes_depois(self):
        from backend.benchmarks.maria_bench.analysis.report import formatar_avisos
        avisos = formatar_avisos(self._resultado(correcoes=[
            {"campo": "nome_arquivo", "antes": "../../teste", "depois": "teste"},
        ]))
        self.assertEqual(len(avisos), 1)
        self.assertIn("corrigido", avisos[0])
        self.assertIn("../../teste", avisos[0])
        self.assertIn("teste", avisos[0])

    def test_json_reparado_gera_aviso(self):
        from backend.benchmarks.maria_bench.analysis.report import formatar_avisos
        avisos = formatar_avisos(self._resultado(
            tool_detected="criar_planilha",
            tool_nome_final="criar_planilha",
            fallbacks=["json_reparado"],
        ))
        self.assertTrue(any("JSON reparado" in a for a in avisos))

    def test_json_limpo_nao_gera_aviso(self):
        from backend.benchmarks.maria_bench.analysis.report import formatar_avisos
        avisos = formatar_avisos(self._resultado(
            tool_detected="criar_planilha",
            tool_nome_final="criar_planilha",
            tool_call_fonte="json",
            fallbacks=[],
        ))
        self.assertEqual(avisos, [])

    def test_chaves_normalizadas_e_colunas_derivadas(self):
        from backend.benchmarks.maria_bench.analysis.report import formatar_avisos
        avisos = formatar_avisos(self._resultado(
            tool_detected="criar_planilha",
            tool_nome_final="criar_planilha",
            fallbacks=["chaves_normalizadas", "colunas_derivadas"],
        ))
        self.assertTrue(any("chaves normalizadas" in a for a in avisos))
        self.assertTrue(any("colunas derivadas" in a for a in avisos))


class TestChatStreamDegeneracao(unittest.TestCase):
    """Testa o abort precoce de geração degenerada em chat_stream."""

    @patch('backend.core.llama_client.requests.Session')
    def test_chat_stream_interrompe_em_degeneracao(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)

        # Reproduz a task 15 (run_20260904_131134): prefixo válido seguido de
        # newlines infinitos — o stream nunca envia finish_reason do servidor.
        linhas = [
            b'data: ' + json.dumps({"choices": [{"delta": {"content": 'criar_documento: ["ata", "Ata", "exemplo: '}, "finish_reason": None}]}).encode(),
        ]
        for _ in range(5):
            linhas.append(b'data: ' + json.dumps({"choices": [{"delta": {"content": "\n" * 50}, "finish_reason": None}]}).encode())

        mock_response = MagicMock(status_code=200)
        mock_response.iter_lines.return_value = iter(linhas)
        mock_session.post.return_value = mock_response

        cliente = LlamaClient()
        metricas = {}
        chunks = list(cliente.chat_stream([{"role": "user", "content": "oi"}], metricas_saida=metricas))

        # Abortou cedo: não consumiu os 250 newlines disponíveis
        textos = "".join(c for c, _ in chunks if c is not None)
        self.assertLess(textos.count("\n"), 100)

        # Sem tool call extraída do lixo; último yield é (None, None)
        ultimo_chunk, ultimo_tool = chunks[-1]
        self.assertIsNone(ultimo_chunk)
        self.assertIsNone(ultimo_tool)

        # Diagnóstico disponível nas métricas
        self.assertTrue(metricas["degeneracao_detectada"])
        self.assertEqual(metricas["finish_reason"], "degenerate")


class TestMariaRunnerDegeneracao(unittest.TestCase):
    """Testa que geração degenerada vira erro descritivo no resultado."""

    def test_degeneracao_gera_erro_descritivo(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask

        class ClienteDegenerado:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                extras = kwargs.get("extras_saida")
                if extras is not None:
                    extras["finish_reason"] = "degenerate"
                    extras["degeneracao_detectada"] = True
                return ("\n" * 120, None, 120, 2.0, 1.0)

        runner = MariaRunner(cliente=ClienteDegenerado())
        task = MariaTask(1, "Teste degeneração", "desc", "Crie um documento de ata com título Ata e conteúdo completo sobre uma reunião.")
        resultado = runner.run(task)

        kinds = [e["kind"] for e in resultado.errors]
        self.assertIn("DegenerateGeneration", kinds)
        self.assertFalse(resultado.runtime_ok)
        self.assertEqual(resultado.finish_reason, "degenerate")


class TestSanitizacaoNomeSeguro(unittest.TestCase):
    """Testa a auto-sanitização silenciosa de nomes inseguros."""

    def test_caminho_relativo_task_24(self):
        from backend.core.tools_schema import _sanitizar_nome_seguro
        self.assertEqual(_sanitizar_nome_seguro("../../teste_seguro"), "teste_seguro")

    def test_caracteres_inseguros_task_25(self):
        from backend.core.tools_schema import _sanitizar_nome_seguro
        # Remove '/' e '*' e strip('.') -> "relatórioseguro" (nome seguro, sem caracteres especiais)
        self.assertEqual(_sanitizar_nome_seguro("../relatório*seguro"), "relatórioseguro")

    def test_nome_valido_inalterado(self):
        from backend.core.tools_schema import _sanitizar_nome_seguro
        self.assertEqual(_sanitizar_nome_seguro("gastos"), "gastos")

    def test_nome_vazio_retorna_placeholder(self):
        from backend.core.tools_schema import _sanitizar_nome_seguro
        self.assertEqual(_sanitizar_nome_seguro(""), "arquivo_sem_nome")


class TestLlamaClientErros(unittest.TestCase):
    """Testa tratamento de erros de conexão e timeout."""


class TestLlamaClientCompatibilidade(unittest.TestCase):
    """Testa métodos de compatibilidade da interface do cliente."""

    @patch('backend.core.llama_client.requests.Session')
    def test_chat_com_tools_stream_compativel(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)

        linhas = [
            b'data: ' + json.dumps({"choices": [{"delta": {"content": "Resposta"}, "finish_reason": None}]}).encode(),
            b'data: ' + json.dumps({"choices": [{"delta": {}, "finish_reason": "stop"}]}).encode(),
        ]
        mock_response = MagicMock(status_code=200)
        mock_response.iter_lines.return_value = iter(linhas)
        mock_session.post.return_value = mock_response

        cliente = LlamaClient()
        chunks = list(cliente.chat_com_tools_stream(
            mensagem_usuario="olá",
            historico=[],
            tools=None,
        ))
        textos = [c for c, _ in chunks if c is not None]
        self.assertIn("Resposta", textos)

    @patch('backend.core.llama_client.requests.Session')
    def test_continuar_com_resultado_ferramenta_stream(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        mock_session.get.return_value = MagicMock(status_code=200)

        linhas = [
            b'data: ' + json.dumps({"choices": [{"delta": {"content": "Feito"}, "finish_reason": None}]}).encode(),
            b'data: ' + json.dumps({"choices": [{"delta": {}, "finish_reason": "stop"}]}).encode(),
        ]
        mock_response = MagicMock(status_code=200)
        mock_response.iter_lines.return_value = iter(linhas)
        mock_session.post.return_value = mock_response

        cliente = LlamaClient()
        metricas = {}
        chunks = list(cliente.continuar_com_resultado_ferramenta_stream(
            historico=[{"role": "system", "content": "sistema"}],
            tool_call={"name": "listar_arquivos", "arguments": {}},
            resultado="pasta vazia",
            tools=None,
            metricas_saida=metricas,
        ))
        textos = [c for c, _ in chunks if c is not None]
        self.assertIn("Feito", textos)
        # Verifica que o payload inclui mensagem role=tool
        payload_enviado = mock_session.post.call_args.kwargs["json"]
        mensagens = payload_enviado["messages"]
        roles = [m["role"] for m in mensagens]
        self.assertIn("tool", roles)

        # Verifica formato correto de tool calling (id, type e tool_call_id)
        msg_assistant = next(m for m in mensagens if m["role"] == "assistant")
        tool_call_payload = msg_assistant["tool_calls"][0]
        self.assertIn("id", tool_call_payload)
        self.assertTrue(tool_call_payload["id"])
        self.assertEqual(tool_call_payload["type"], "function")
        self.assertEqual(
            tool_call_payload["function"]["name"], "listar_arquivos"
        )
        self.assertEqual(
            tool_call_payload["function"]["arguments"],
            json.dumps({}, ensure_ascii=False),
        )
        msg_tool = next(m for m in mensagens if m["role"] == "tool")
        self.assertEqual(msg_tool["tool_call_id"], tool_call_payload["id"])


class TestSegurancaComandosBridge(unittest.TestCase):
    """Testes de segurança para upload_arquivo e transcrever_audio (P0)."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_fora = tempfile.TemporaryDirectory()
        self.original_pastas = os.environ.get("PASTAS_PERMITIDAS")
        os.environ["PASTAS_PERMITIDAS"] = self.temp_dir.name
        self.pasta_gerados = os.path.join(self.temp_dir.name, "arquivos_gerados")
        self.original_gerados = os.environ.get("PASTA_ARQUIVOS_GERADOS")
        os.environ["PASTA_ARQUIVOS_GERADOS"] = self.pasta_gerados

    def tearDown(self):
        self.temp_dir.cleanup()
        self.temp_fora.cleanup()
        if self.original_pastas:
            os.environ["PASTAS_PERMITIDAS"] = self.original_pastas
        else:
            os.environ.pop("PASTAS_PERMITIDAS", None)
        if self.original_gerados:
            os.environ["PASTA_ARQUIVOS_GERADOS"] = self.original_gerados
        else:
            os.environ.pop("PASTA_ARQUIVOS_GERADOS", None)
        os.environ.pop("WHISPER_BIN", None)

    def _despachar(self, comando, payload):
        from backend.main import _despachar_comando
        return _despachar_comando(None, comando, payload)

    def test_upload_rejeita_diretorio(self):
        """upload_arquivo deve rejeitar pastas (apenas arquivos)."""
        status, dados, erro = self._despachar(
            "upload_arquivo", {"caminho": self.temp_fora.name}
        )
        self.assertEqual(status, "erro")
        self.assertIn("inválido", erro)

    def test_upload_arquivo_valido_copia_para_pasta_gerenciada(self):
        """upload_arquivo copia um arquivo válido para a pasta permitida."""
        origem = os.path.join(self.temp_fora.name, "doc.txt")
        with open(origem, "w", encoding="utf-8") as f:
            f.write("conteúdo")

        status, dados, erro = self._despachar("upload_arquivo", {"caminho": origem})

        self.assertEqual(status, "ok", erro)
        self.assertTrue(
            os.path.isfile(os.path.join(self.pasta_gerados, "doc.txt")),
            f"Arquivo deveria estar em {self.pasta_gerados}: {dados}",
        )

    def test_transcrever_audio_rejeita_caminho_fora_das_pastas_permitidas(self):
        """transcrever_audio não pode ler/deletar arquivos fora das pastas
        permitidas (corrige deleção arbitrária via unlink)."""
        audio_fora = os.path.join(self.temp_fora.name, "audio.wav")
        with open(audio_fora, "wb") as f:
            f.write(b"RIFF....")

        status, dados, erro = self._despachar(
            "transcrever_audio", {"caminho": audio_fora}
        )

        self.assertEqual(status, "erro")
        # O arquivo fora da pasta permitida NÃO pode ser deletado
        self.assertTrue(os.path.exists(audio_fora))

    def test_transcrever_audio_rejeita_whisper_bin_invalido(self):
        """WHISPER_BIN com caminho/espaços deve ser rejeitado (evita
        execução de binário arbitrário)."""
        audio = os.path.join(self.temp_dir.name, "audio.wav")
        with open(audio, "wb") as f:
            f.write(b"RIFF....")
        os.environ["WHISPER_BIN"] = "C:/evil/dir/whisper-main.exe"

        status, dados, erro = self._despachar(
            "transcrever_audio", {"caminho": audio}
        )

        self.assertEqual(status, "erro")
        self.assertIn("WHISPER_BIN", erro)


class TestSegurancaApiHttp(unittest.TestCase):
    """Testes de autenticação da API bridge HTTP (P1)."""

    def setUp(self):
        from backend.main import _criar_app_http, _carregar_token_api
        self.token = _carregar_token_api()
        self.app = _criar_app_http(None, self.token)
        self.client = self.app.test_client()

    def test_chat_sem_token_rejeitado(self):
        resp = self.client.post("/chat", json={"id": "1", "comando": "status", "dados": {}})
        self.assertEqual(resp.status_code, 401)

    def test_chat_token_invalido_rejeitado(self):
        resp = self.client.post(
            "/chat",
            json={"id": "1", "comando": "status", "dados": {}},
            headers={"Authorization": "Bearer token-falso"},
        )
        self.assertEqual(resp.status_code, 401)

    def test_ping_fica_aberto_para_health_check(self):
        resp = self.client.get("/ping")
        self.assertEqual(resp.status_code, 200)

    def test_cors_origem_maliciosa_nao_autorizada(self):
        resp = self.client.options(
            "/chat",
            headers={
                "Origin": "https://site-malicioso.example",
                "Access-Control-Request-Method": "POST",
            },
        )
        self.assertNotEqual(
            resp.headers.get("Access-Control-Allow-Origin", ""), "*"
        )

    def test_cors_origem_tauri_autorizada(self):
        resp = self.client.options(
            "/chat",
            headers={
                "Origin": "http://tauri.localhost",
                "Access-Control-Request-Method": "POST",
                "Access-Control-Request-Headers": "authorization",
            },
        )
        self.assertEqual(
            resp.headers.get("Access-Control-Allow-Origin"), "http://tauri.localhost"
        )


def _fts5_disponivel() -> bool:
    """Verifica em tempo de execução se o SQLite local suporta FTS5."""
    import sqlite3
    try:
        conexao_teste = sqlite3.connect(":memory:")
        conexao_teste.execute("CREATE VIRTUAL TABLE t USING fts5(x)")
        conexao_teste.close()
        return True
    except sqlite3.OperationalError:
        return False


class TestManualRedacaoIngestao(unittest.TestCase):
    """Testes para o parser/chunker do Manual de Redação (ingest_manual_redacao.py)."""

    def test_extrair_chunks_classifica_secao_oficio(self):
        from backend.database.ingest_manual_redacao import extrair_chunks
        texto = (
            "#### **5 O padrão ofício**\n"
            "Texto introdutório do padrão ofício.\n"
            "##### 5.1.7 Fechos para comunicações\n"
            "Conteúdo bastante longo sobre os fechos oficiais utilizados em comunicações formais do governo federal.\n"
            "##### 5.1.8 Identificação do signatário\n"
            "Conteúdo sobre como identificar corretamente o signatário do documento oficial em questão.\n"
        )
        chunks = extrair_chunks(texto)
        secoes = {c["secao"]: c["tipo_documento"] for c in chunks}
        self.assertIn("5.1.7 Fechos para comunicações", secoes)
        self.assertEqual(secoes["5.1.7 Fechos para comunicações"], "oficio")

    def test_extrair_chunks_descarta_secoes_vazias(self):
        from backend.database.ingest_manual_redacao import extrair_chunks
        texto = "# **Capítulo I**\n\n#### **1 Panorama**\num\n"
        chunks = extrair_chunks(texto)
        self.assertEqual(chunks, [])

    def test_classificar_tipo_documento_mapeia_exposicao_e_mensagem(self):
        from backend.database.ingest_manual_redacao import _classificar_tipo_documento
        self.assertEqual(_classificar_tipo_documento("6.2.1 Definição e finalidade"), "exposicao_motivos")
        self.assertEqual(_classificar_tipo_documento("6.3.2 Forma e estrutura"), "mensagem")
        self.assertEqual(_classificar_tipo_documento("6.4.1 Definição e finalidade"), "email")
        self.assertEqual(_classificar_tipo_documento("Prefácio"), "geral")


@unittest.skipUnless(_fts5_disponivel(), "SQLite local sem suporte a FTS5")
class TestManualRedacaoConsulta(unittest.TestCase):
    """Testes para o módulo de consulta RAG (core/manual_redacao.py)."""

    def setUp(self):
        import backend.database.connection as connection_module
        connection_module._CONNECTION = None
        self.temp_dir = tempfile.TemporaryDirectory()
        connection_module.init_db(f"{self.temp_dir.name}/teste_manual.db")
        from backend.database.schema import init_db
        init_db()
        conn = connection_module.get_connection()
        conn.executemany(
            "INSERT INTO manual_redacao_fts (tipo_documento, secao, conteudo) VALUES (?, ?, ?)",
            [
                ("oficio", "5.1.7 Fechos para comunicações",
                 "O fecho das comunicações oficiais objetiva saudar o destinatário com respeitosamente ou atenciosamente."),
                ("oficio", "4.4 Vocativo",
                 "O vocativo a ser empregado em comunicações dirigidas aos chefes de poder é Senhor mais o cargo."),
                ("mensagem", "6.3.2 Forma e estrutura",
                 "A mensagem não obedece a um único e rígido modelo, mas apresenta a mesma estrutura do padrão ofício."),
            ],
        )
        conn.commit()

    def tearDown(self):
        import backend.database.connection as connection_module
        connection_module.close_connection()
        connection_module._DB_PATH = None  # evita que _DB_PATH aponte para o temp_dir já removido
        self.temp_dir.cleanup()

    def test_consulta_com_tipo_e_termo_busca(self):
        from backend.core.manual_redacao import consultar_manual
        resultado = consultar_manual(tipo_documento="oficio", termo_busca="vocativo")
        self.assertIn("Vocativo", resultado)
        self.assertNotIn("6.3.2", resultado)

    def test_consulta_sem_termo_busca_filtra_apenas_por_tipo(self):
        from backend.core.manual_redacao import consultar_manual
        resultado = consultar_manual(tipo_documento="mensagem")
        self.assertIn("6.3.2 Forma e estrutura", resultado)

    def test_consulta_sem_resultados_retorna_mensagem_amigavel(self):
        from backend.core.manual_redacao import consultar_manual
        resultado = consultar_manual(tipo_documento="email", termo_busca="anexo")
        self.assertIn("Nenhum trecho", resultado)

    def test_consulta_tipo_invalido_e_ignorado_sem_erro(self):
        from backend.core.manual_redacao import consultar_manual
        resultado = consultar_manual(tipo_documento="tipo_que_nao_existe", termo_busca="fecho")
        self.assertIn("Fechos", resultado)

    def test_consulta_trunca_trechos_muito_longos(self):
        """Cobre a correção da Tarefa 3: trechos maiores que
        MANUAL_REDACAO_MAX_CHARS_POR_TRECHO devem ser truncados com
        marcador '[...]' para não estourar o contexto do modelo."""
        from backend.core.manual_redacao import consultar_manual
        import backend.database.connection as connection_module

        conn = connection_module.get_connection()
        conteudo_longo = "palavra " * 500  # muito acima do limite padrão (800 caracteres)
        conn.execute(
            "INSERT INTO manual_redacao_fts (tipo_documento, secao, conteudo) VALUES (?, ?, ?)",
            ("geral", "99.9 Seção de teste longa", conteudo_longo),
        )
        conn.commit()

        resultado = consultar_manual(tipo_documento="geral", termo_busca="palavra")
        self.assertIn("[...]", resultado)


class TestFerramentaConsultarManualRedacao(unittest.TestCase):
    """Testes para a integração da ferramenta em tools_schema.py."""

    def test_consultar_manual_redacao_esta_em_ferramentas_leitura(self):
        self.assertIn("consultar_manual_redacao", FERRAMENTAS_LEITURA)

    def test_ferramenta_criar_documento_nao_menciona_memorandos(self):
        """Cobre a correção de consistência da Tarefa 4.5: 'memorando' foi
        unificado sob 'ofício' no Manual, então criar_documento não deve
        mais listá-lo como exemplo de uso direto."""
        descricao = FERRAMENTA_CRIAR_DOCUMENTO["function"]["description"]
        self.assertNotIn("memorandos", descricao)

    def test_executar_ferramenta_leitura_ferramenta_desconhecida_ainda_falha(self):
        with self.assertRaises(ValueError):
            executar_ferramenta_leitura("ferramenta_inexistente", {})

class TestObterMetadadosModelo(unittest.TestCase):
    """Testa a extração de metadados do llama-server via /v1/models (mock, sem servidor)."""

    def test_obter_metadados_resposta_valida(self):
        from backend.benchmarks.maria_bench.run_benchmark import _obter_metadados_modelo
        mock_response = MagicMock(status_code=200)
        mock_response.json.return_value = {
            "data": [{
                "id": "C:\\blob\\sha256-2bada8a",
                "meta": {"n_params": 3397103616, "n_vocab": 151936,
                         "n_ctx": 8192, "n_ctx_train": 131072,
                         "size": 2098976768, "ftype": 14},
            }]
        }
        with patch("backend.benchmarks.maria_bench.run_benchmark._requests.get", return_value=mock_response):
            m = _obter_metadados_modelo()
        self.assertIsNotNone(m)
        self.assertEqual(m["quantizacao"], "Q4_K_M")
        self.assertEqual(m["rotulo_tamanho"], "Qwen2.5 3B")
        self.assertEqual(m["id_exibicao"], "Qwen2.5 3B")

    def test_obter_metadados_servidor_offline(self):
        import requests as req
        from backend.benchmarks.maria_bench.run_benchmark import _obter_metadados_modelo

        with patch("backend.benchmarks.maria_bench.run_benchmark._requests.get",
                   side_effect=req.exceptions.ConnectionError()):
            self.assertIsNone(_obter_metadados_modelo())

    def test_obter_metadados_status_500(self):
        from backend.benchmarks.maria_bench.run_benchmark import _obter_metadados_modelo

        mock_response = MagicMock(status_code=500)
        with patch("backend.benchmarks.maria_bench.run_benchmark._requests.get", return_value=mock_response):
            self.assertIsNone(_obter_metadados_modelo())


class TestDerivarRotuloModelo(unittest.TestCase):
    """Testa a derivação de rótulo legível a partir de n_params e n_vocab."""

    def test_qwen25_3b(self):
        from backend.benchmarks.maria_bench.run_benchmark import _derivar_rotulo_modelo
        self.assertEqual(_derivar_rotulo_modelo(3397103616, 151936), "Qwen2.5 3B")

    def test_qwen25_7b(self):
        from backend.benchmarks.maria_bench.run_benchmark import _derivar_rotulo_modelo
        self.assertEqual(_derivar_rotulo_modelo(7615616512, 151936), "Qwen2.5 7B")

    def test_qwen25_14b(self):
        from backend.benchmarks.maria_bench.run_benchmark import _derivar_rotulo_modelo
        self.assertEqual(_derivar_rotulo_modelo(14771111936, 151936), "Qwen2.5 14B")

    def test_vocab_desconhecido(self):
        from backend.benchmarks.maria_bench.run_benchmark import _derivar_rotulo_modelo
        self.assertEqual(_derivar_rotulo_modelo(3397103616, 99999), "3B")

    def test_sem_params(self):
        from backend.benchmarks.maria_bench.run_benchmark import _derivar_rotulo_modelo
        self.assertEqual(_derivar_rotulo_modelo(0, 151936), "")


class TestPareceCaminhoLocal(unittest.TestCase):
    """Testa a detecção de caminhos locais/blobs."""

    def test_blob_local(self):
        from backend.benchmarks.maria_bench.run_benchmark import _parece_caminho_local
        self.assertTrue(_parece_caminho_local("C:\\blob\\sha256-2bada8a"))

    def test_hash_puro(self):
        from backend.benchmarks.maria_bench.run_benchmark import _parece_caminho_local
        self.assertTrue(_parece_caminho_local("2bada8a7450677000f678be90653b85d364de7db25eb5ea54136ada5f3933730"))

    def test_gguf(self):
        from backend.benchmarks.maria_bench.run_benchmark import _parece_caminho_local
        self.assertTrue(_parece_caminho_local("models/qwen2.gguf"))

    def test_nome_simples(self):
        from backend.benchmarks.maria_bench.run_benchmark import _parece_caminho_local
        self.assertFalse(_parece_caminho_local("qwen2.5-omni-3b"))

    def test_vazio(self):
        from backend.benchmarks.maria_bench.run_benchmark import _parece_caminho_local
        self.assertFalse(_parece_caminho_local(""))
        self.assertFalse(_parece_caminho_local(None))


class TestAlertaNaoDisparaParaBlob(unittest.TestCase):
    """Testa que o alerta de divergência NÃO dispara quando o id é blob
    do mesmo modelo configurado."""

    def test_sem_alerta_para_blob_mesmo_modelo(self):
        import io
        from contextlib import redirect_stdout
        from backend.benchmarks.maria_bench.run_benchmark import _warmup_model

        metadados = {
            "id": "C:\\blob\\sha256-2bada8a",
            "id_exibicao": "Qwen2.5 3B",
            "rotulo_tamanho": "Qwen2.5 3B",
            "n_ctx": 2048,
        }

        class FakeClient:
            def enviar_mensagem(self, *a, **kw):
                return "ok"

        buf = io.StringIO()
        with patch("backend.benchmarks.maria_bench.run_benchmark._obter_metadados_modelo", return_value=metadados), \
             patch("backend.benchmarks.maria_bench.run_benchmark._contar_tokens_exatos", return_value=820), \
             patch("backend.benchmarks.maria_bench.run_benchmark.LlamaClient", return_value=FakeClient()):
            with redirect_stdout(buf):
                meta = _warmup_model()

        # "3b" está contido em "qwen2.5-omni-3b" -> NÃO deve gerar alerta
        self.assertNotIn("[AVISO]", buf.getvalue())


class TestAvisoNctx(unittest.TestCase):
    """Testa o aviso de n_ctx quando config > servidor."""

    def test_relatorio_contem_aviso_nctx(self):
        import tempfile
        from unittest.mock import patch as _patch
        from backend.benchmarks.maria_bench.analysis.report import generate_report
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTaskResult

        results = [
            MariaTaskResult(
                task_id=1, task_name="T1", category="conversa", model="m",
                tool_detected=None, tool_correct=True, confirmation_completed=True,
                keyword_match=True, runtime_ok=True, final_message="ok",
                latency_ms=100.0, errors=[], raw_tool_args={},
            ),
        ]

        metadados = {
            "id": "modelo.gguf", "id_exibicao": "modelo.gguf",
            "n_ctx": 4096, "n_ctx_train": 131072,
            "quantizacao": "Q4_K_M", "n_params": 3397103616,
            "tamanho_legivel": "1.95 GiB",
        }

        with _patch("backend.benchmarks.maria_bench.analysis.report.LLAMA_NUM_CTX", 8192):
            with tempfile.TemporaryDirectory() as tmpdir:
                generate_report(
                    results,
                    MagicMock(
                        total_tasks=1, tool_accuracy=1.0, confirmation_success_rate=1.0,
                        keyword_match_rate=1.0, runtime_success_rate=1.0,
                        language_compliance_rate=1.0, args_accuracy=1.0,
                        avg_tokens_por_segundo=0.0, avg_ttft_ms=None,
                        p50_latency_ms=100.0, p90_latency_ms=100.0,
                        avg_latency_ms=100.0, error_distribution={}, by_category={},
                        contexto_ok_rate=1.0,
                        semantic_quality_rate=1.0, semantic_errors_by_type={},
                        correcoes_count=0,
                    ),
                    tmpdir,
                    metadados_modelo=metadados,
                )
                with open(os.path.join(tmpdir, "report.md"), encoding="utf-8") as f:
                    report = f.read()

        self.assertIn("4096", report)
        self.assertIn("n_ctx", report)


class TestFtypeParaNome(unittest.TestCase):
    """Testa a conversão de ftype (enum GGML) para nome legível."""

    def test_q4_k_m(self):
        from backend.benchmarks.maria_bench.run_benchmark import _ftype_para_nome
        self.assertEqual(_ftype_para_nome(14), "Q4_K_M")

    def test_string_passa_direto(self):
        from backend.benchmarks.maria_bench.run_benchmark import _ftype_para_nome
        self.assertEqual(_ftype_para_nome("Q8_0"), "Q8_0")

    def test_none(self):
        from backend.benchmarks.maria_bench.run_benchmark import _ftype_para_nome
        self.assertEqual(_ftype_para_nome(None), "")


class TestSamplerParamsBenchmark(unittest.TestCase):
    """Testa os parâmetros de sampler configuráveis e sua exposição no benchmark."""

    def test_config_defaults_dos_sampler_params(self):
        from backend.core.config import (
            LLAMA_DRY_ALLOWED_LENGTH, LLAMA_DRY_BASE, LLAMA_DRY_MULTIPLIER,
            LLAMA_DRY_PENALTY_LAST_N, LLAMA_FREQUENCY_PENALTY, LLAMA_MIN_P,
            LLAMA_PRESENCE_PENALTY, LLAMA_REPEAT_LAST_N, LLAMA_REPEAT_PENALTY,
            LLAMA_TOP_K, LLAMA_TOP_N_SIGMA, LLAMA_TOP_P, LLAMA_TYPICAL_P,
            LLAMA_XTC_PROBABILITY, LLAMA_XTC_THRESHOLD,
        )
        self.assertEqual(LLAMA_REPEAT_LAST_N, 128)
        self.assertEqual(LLAMA_REPEAT_PENALTY, 1.1)  # 1.1: default clássico; 1.3 quebra tool calling (ver config.py)
        self.assertEqual(LLAMA_FREQUENCY_PENALTY, 0.0)
        self.assertEqual(LLAMA_PRESENCE_PENALTY, 0.0)
        self.assertEqual(LLAMA_DRY_MULTIPLIER, 0.8)
        self.assertEqual(LLAMA_DRY_BASE, 1.75)
        self.assertEqual(LLAMA_DRY_ALLOWED_LENGTH, 2)
        self.assertEqual(LLAMA_DRY_PENALTY_LAST_N, 64)
        self.assertEqual(LLAMA_TOP_K, 40)
        self.assertEqual(LLAMA_TOP_P, 0.95)
        self.assertEqual(LLAMA_MIN_P, 0.05)
        self.assertEqual(LLAMA_XTC_PROBABILITY, 0.0)
        self.assertEqual(LLAMA_XTC_THRESHOLD, 0.1)
        self.assertEqual(LLAMA_TYPICAL_P, 1.0)
        self.assertEqual(LLAMA_TOP_N_SIGMA, -1.0)

    def test_montar_sampler_params_contem_todos_os_campos(self):
        from backend.core.llama_client import montar_sampler_params
        params = montar_sampler_params()
        esperados = {
            "temperature", "repeat_last_n", "repeat_penalty",
            "frequency_penalty", "presence_penalty", "dry_multiplier",
            "dry_base", "dry_allowed_length", "dry_penalty_last_n",
            "top_k", "top_p", "min_p", "xtc_probability",
            "xtc_threshold", "typical_p", "top_n_sigma",
        }
        self.assertEqual(set(params), esperados)
        self.assertEqual(params["top_k"], 40)
        self.assertEqual(params["top_p"], 0.95)
        self.assertEqual(params["temperature"], 0.1)

    def test_montar_payload_inclui_sampler_params_com_tools(self):
        from backend.core.llama_client import LlamaClient
        cliente = LlamaClient(model="modelo-teste")
        payload = cliente._montar_payload(
            [{"role": "user", "content": "oi"}], tools=[{}], stream=False,
            incluir_temperatura=True,
        )
        for chave in ("temperature", "top_k", "top_p", "min_p",
                      "repeat_last_n", "repeat_penalty", "dry_base",
                      "xtc_probability", "typical_p", "top_n_sigma"):
            self.assertIn(chave, payload)
        self.assertEqual(payload["top_k"], 40)
        self.assertEqual(payload["temperature"], 0.1)

    def test_montar_payload_omite_sampler_sem_tools(self):
        from backend.core.llama_client import LlamaClient
        cliente = LlamaClient(model="modelo-teste")
        payload = cliente._montar_payload(
            [{"role": "user", "content": "oi"}], tools=None, stream=False,
        )
        self.assertNotIn("top_k", payload)
        self.assertNotIn("temperature", payload)

    def test_maria_task_result_campos_novos_com_default(self):
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTaskResult
        result = MariaTaskResult(
            task_id=1, task_name="T", category="conversa", model="m",
            tool_detected=None, tool_correct=True, confirmation_completed=True,
            keyword_match=True, runtime_ok=True, final_message="ok",
            latency_ms=100.0, errors=[], raw_tool_args={},
        )
        self.assertEqual(result.prompt_enviado, [])
        self.assertEqual(result.resposta_bruta_modelo, "")
        self.assertEqual(result.sampler_params, {})

    def test_runner_preenche_prompt_e_resposta_bruta(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteFalso:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return (
                    "Vou criar a planilha.",
                    {"name": "criar_planilha", "arguments": {
                        "nome_arquivo": "gastos", "colunas": ["Data", "Valor"]}},
                    8, 5.0, 1.0,
                )

        task = MariaTask(
            9101, "Prompt/resposta", "desc", "Crie uma planilha de gastos.",
            expected_tool="criar_planilha", confirm_sequence=["sim"],
            category=MariaTaskCategory.CRIAR_PLANILHA,
        )
        runner = MariaRunner(cliente=ClienteFalso())
        resultado = runner.run(task)

        # Prompt enviado contém system (reforço) e a mensagem do usuário.
        self.assertTrue(any(m["role"] == "system" for m in resultado.prompt_enviado))
        self.assertEqual(resultado.prompt_enviado[-1]["role"], "user")
        self.assertIn("Crie uma planilha de gastos.", resultado.prompt_enviado[-1]["content"])
        # Resposta bruta preserva o que o modelo gerou ANTES da confirmação.
        self.assertEqual(resultado.resposta_bruta_modelo, "Vou criar a planilha.")
        self.assertNotEqual(resultado.final_message, resultado.resposta_bruta_modelo)
        self.assertTrue(resultado.sampler_params)
        self.assertIn("top_k", resultado.sampler_params)

    def test_report_contem_parametros_e_detalhes_por_execucao(self):
        import tempfile
        from unittest.mock import MagicMock
        from backend.benchmarks.maria_bench.analysis.report import generate_report
        from backend.core.llama_client import montar_sampler_params
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTaskResult

        results = [
            MariaTaskResult(
                task_id=3, task_name="Planilha básica", category="criar_planilha",
                model="m", tool_detected="criar_planilha", tool_correct=True,
                confirmation_completed=True, keyword_match=True, runtime_ok=True,
                final_message="Planilha criada com sucesso: ...",
                latency_ms=100.0, errors=[], raw_tool_args={},
                prompt_enviado=[
                    {"role": "system", "content": "system"},
                    {"role": "user", "content": "Crie uma planilha de gastos"},
                ],
                resposta_bruta_modelo="Vou criar a planilha.",
                sampler_params=montar_sampler_params(),
            ),
        ]
        metrics = MagicMock(
            total_tasks=1, tool_accuracy=1.0, confirmation_success_rate=1.0,
            keyword_match_rate=1.0, runtime_success_rate=1.0,
            language_compliance_rate=1.0, args_accuracy=1.0,
            avg_tokens_por_segundo=0.0, avg_ttft_ms=None,
            p50_latency_ms=100.0, p90_latency_ms=100.0,
            avg_latency_ms=100.0, error_distribution={}, by_category={},
            contexto_ok_rate=1.0,
            semantic_quality_rate=1.0, semantic_errors_by_type={},
            correcoes_count=0,
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            generate_report(
                results, metrics, tmpdir, sampler_params=montar_sampler_params(),
                detail=True,
            )
            with open(os.path.join(tmpdir, "report.md"), encoding="utf-8") as f:
                report = f.read()

        self.assertIn("## Parâmetros do sampler", report)
        self.assertIn("| top_k | 40 |", report)
        self.assertIn("| temperature | 0.100 |", report)
        self.assertIn("## Detalhes por execução", report)
        self.assertIn("Crie uma planilha de gastos", report)
        self.assertIn("Vou criar a planilha.", report)


class TestExtrairNomeExibicao(unittest.TestCase):
    """Testa a extração de nome amigável a partir do ID cru do modelo."""

    def test_id_completo(self):
        from backend.benchmarks.maria_bench.run_benchmark import _extrair_nome_exibicao
        self.assertEqual(
            _extrair_nome_exibicao("ggml-org/Qwen2.5-Omni-3B-GGUF:Q4_K_M"),
            "Qwen2.5 Omni 3B",
        )

    def test_id_simples_com_hifens(self):
        from backend.benchmarks.maria_bench.run_benchmark import _extrair_nome_exibicao
        self.assertEqual(_extrair_nome_exibicao("qwen2.5-omni-3b"), "qwen2.5 omni 3b")

    def test_id_com_underscores(self):
        from backend.benchmarks.maria_bench.run_benchmark import _extrair_nome_exibicao
        self.assertEqual(_extrair_nome_exibicao("Qwen2.5_Omni_3B"), "Qwen2.5 Omni 3B")

    def test_id_vazio(self):
        from backend.benchmarks.maria_bench.run_benchmark import _extrair_nome_exibicao
        self.assertEqual(_extrair_nome_exibicao(""), "")


class TestExtrairQuantizacao(unittest.TestCase):
    """Testa a extração da quantização a partir do ID do modelo."""

    def test_extrai_do_sufixo(self):
        from backend.benchmarks.maria_bench.run_benchmark import _extrair_quantizacao
        self.assertEqual(
            _extrair_quantizacao("ggml-org/Qwen2.5-Omni-3B-GGUF:Q4_K_M"), "Q4_K_M"
        )

    def test_sem_sufixo_retorna_desconhecida(self):
        from backend.benchmarks.maria_bench.run_benchmark import _extrair_quantizacao
        self.assertEqual(_extrair_quantizacao("qwen2.5-omni-3b"), "desconhecida")


class TestContextoOk(unittest.TestCase):
    """Testa o campo contexto_ok no resultado e nas métricas do benchmark."""

    def test_default_true_no_resultado(self):
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTaskResult
        result = MariaTaskResult(
            task_id=1, task_name="T", category="conversa", model="m",
            tool_detected=None, tool_correct=True, confirmation_completed=True,
            keyword_match=True, runtime_ok=True, final_message="ok",
            latency_ms=100.0, errors=[], raw_tool_args={},
        )
        self.assertTrue(result.contexto_ok)

    def test_contexto_ok_rate_nas_metricas(self):
        from backend.benchmarks.maria_bench.analysis.metrics import calculate_maria_metrics
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTaskResult
        base = dict(
            task_name="T", category="conversa", model="m",
            tool_detected=None, tool_correct=True, confirmation_completed=True,
            keyword_match=True, runtime_ok=True, final_message="ok",
            latency_ms=100.0, errors=[], raw_tool_args={},
        )
        r_ok = MariaTaskResult(**base, task_id=1, contexto_ok=True)
        r_falha = MariaTaskResult(**base, task_id=2, contexto_ok=False)
        metrics = calculate_maria_metrics([r_ok, r_falha])
        self.assertEqual(metrics.contexto_ok_rate, 0.5)

    def test_runner_detecta_erro_de_contexto(self):
        from unittest.mock import patch as _patch
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner, LlamaClientError
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteErroContexto:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                # Mesma classe que o runner captura (import de core.llama_client).
                raise LlamaClientError(
                    "Prompt token count 5000 exceeds the available context size 4096"
                )

        task = MariaTask(
            9107, "Contexto", "desc", "Crie uma planilha grande.",
            category=MariaTaskCategory.CRIAR_PLANILHA,
        )
        with _patch("backend.benchmarks.maria_bench.runners.maria_runner.time.sleep"):
            resultado = MariaRunner(cliente=ClienteErroContexto()).run(task)

        self.assertFalse(resultado.contexto_ok)
        self.assertTrue(
            any(e.get("kind") == "LlamaClientError" for e in resultado.errors)
        )

    def test_runner_erro_generico_mantem_contexto_ok(self):
        from unittest.mock import patch as _patch
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner, LlamaClientError
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteErroRede:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                raise LlamaClientError("Falha de conexao com o llama-server")

        task = MariaTask(
            9108, "Erro rede", "desc", "Crie uma planilha.",
            category=MariaTaskCategory.CRIAR_PLANILHA,
        )
        with _patch("backend.benchmarks.maria_bench.runners.maria_runner.time.sleep"):
            resultado = MariaRunner(cliente=ClienteErroRede()).run(task)

        self.assertTrue(resultado.contexto_ok)
        self.assertTrue(
            any(e.get("kind") == "LlamaClientError" for e in resultado.errors)
        )


class TestEstimarTokens(unittest.TestCase):
    """Testa a estimativa rápida de tokens (~4 caracteres por token)."""

    def test_texto_vazio_retorna_zero(self):
        from backend.benchmarks.maria_bench.utils import estimar_tokens
        self.assertEqual(estimar_tokens(""), 0)

    def test_400_caracteres_retorna_100(self):
        from backend.benchmarks.maria_bench.utils import estimar_tokens
        self.assertEqual(estimar_tokens("a" * 400), 100)

    def test_texto_curto_retorna_no_minimo_um(self):
        from backend.benchmarks.maria_bench.utils import estimar_tokens
        self.assertEqual(estimar_tokens("ok"), 1)


class TestCalibracaoDeTokens(unittest.TestCase):
    """Testa o fator de calibração medido no warmup via /tokenize."""

    def setUp(self):
        # O warmup de outros testes define o fator global; cada teste parte de 1.0.
        from backend.benchmarks.maria_bench import utils
        utils._fator_calibracao = 1.0

    def tearDown(self):
        from backend.benchmarks.maria_bench import utils
        utils._fator_calibracao = 1.0

    def test_fator_medido_e_aplicado(self):
        from backend.benchmarks.maria_bench import utils
        fator = utils.definir_fator_calibracao(800, "a" * 400)
        self.assertEqual(fator, 8.0)
        self.assertEqual(utils.estimar_tokens_calibrado("a" * 400), 800)

    def test_sem_calibracao_estimativa_pura(self):
        from backend.benchmarks.maria_bench import utils
        self.assertEqual(utils.obter_fator_calibracao(), 1.0)
        self.assertEqual(utils.estimar_tokens_calibrado("a" * 400), 100)

    def test_texto_vazio_retorna_zero_mesmo_com_fator(self):
        from backend.benchmarks.maria_bench import utils
        utils.definir_fator_calibracao(800, "a" * 400)
        self.assertEqual(utils.estimar_tokens_calibrado(""), 0)

    def test_fator_ignorado_se_estimativa_zero(self):
        from backend.benchmarks.maria_bench import utils
        utils.definir_fator_calibracao(800, "")
        self.assertEqual(utils.obter_fator_calibracao(), 1.0)


class TestWarmupCtxSize(unittest.TestCase):
    """Testa a detecção de contexto real e a verificação do system prompt no warmup."""

    def _executar_warmup(self, metadados, contar_tokens):
        import io
        from contextlib import redirect_stdout
        from backend.benchmarks.maria_bench import run_benchmark

        class FakeClient:
            def enviar_mensagem(self, *a, **kw):
                return "ok"

        buf = io.StringIO()
        with patch("backend.benchmarks.maria_bench.run_benchmark._obter_metadados_modelo", return_value=metadados), \
             patch("backend.benchmarks.maria_bench.run_benchmark._contar_tokens_exatos", side_effect=contar_tokens), \
             patch("backend.benchmarks.maria_bench.run_benchmark.LlamaClient", return_value=FakeClient()):
            with redirect_stdout(buf):
                meta = run_benchmark._warmup_model()
        return meta, buf.getvalue()

    def test_ctx_real_via_models(self):
        metadados = {"id": "m", "id_exibicao": "m", "n_ctx": 2048}
        meta, _ = self._executar_warmup(metadados, lambda *a: 820)
        self.assertEqual(meta["ctx_size"], 2048)
        self.assertEqual(meta["ctx_fonte"], "models")

    def test_fallback_para_llama_num_ctx_com_aviso(self):
        from backend.benchmarks.maria_bench.run_benchmark import LLAMA_NUM_CTX
        metadados = {"id": "m", "id_exibicao": "m"}  # sem n_ctx
        meta, saida = self._executar_warmup(metadados, lambda *a: None)
        self.assertEqual(meta["ctx_size"], LLAMA_NUM_CTX)
        self.assertEqual(meta["ctx_fonte"], "fallback")
        self.assertIn("[AVISO]", saida)

    def test_system_prompt_tokens_contagem_exata(self):
        metadados = {"id": "m", "id_exibicao": "m", "n_ctx": 2048}
        meta, _ = self._executar_warmup(metadados, lambda *a: 820)
        self.assertEqual(meta["system_prompt_tokens"], 820)

    def test_aborta_quando_prompt_nao_cabe(self):
        metadados = {"id": "m", "id_exibicao": "m", "n_ctx": 2048}
        with self.assertRaises(SystemExit):
            self._executar_warmup(metadados, lambda *a: 100000)

    def tearDown(self):
        # Evita vazar o fator de calibração global para outros testes.
        from backend.benchmarks.maria_bench import utils
        utils._fator_calibracao = 1.0


class TestNumCtxAdaptativo(unittest.TestCase):
    """Testa a remoção adaptativa de num_ctx quando o servidor rejeita (HTTP 400)."""

    @patch('backend.core.llama_client.requests.Session')
    def test_400_remove_num_ctx_e_refaz_sem_o_campo(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        resp_models = MagicMock()
        resp_models.status_code = 200
        mock_session.get.return_value = resp_models

        resp_400 = MagicMock()
        resp_400.status_code = 400
        resp_400.text = "erro: num_ctx invalido"
        resp_200 = MagicMock()
        resp_200.status_code = 200
        # Captura o payload NO MOMENTO da chamada: _make_request muta o mesmo
        # dict (pop de num_ctx), então call_args_list guardaria a versão final.
        respostas = [resp_400, resp_200]
        payloads_chamados = []

        def capturar_post(*args, **kwargs):
            payloads_chamados.append(dict(kwargs.get("json") or {}))
            return respostas.pop(0)

        mock_session.post.side_effect = capturar_post

        cliente = LlamaClient()
        payload = {"model": "m", "messages": [], "num_ctx": 8192}
        response = cliente._make_request(payload)

        self.assertIs(response, resp_200)
        self.assertFalse(cliente._num_ctx_respeitado)
        self.assertIn("num_ctx", payloads_chamados[0])
        self.assertNotIn("num_ctx", payloads_chamados[1])

    @patch('backend.core.llama_client.requests.Session')
    def test_200_mantem_num_ctx(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        resp_models = MagicMock()
        resp_models.status_code = 200
        mock_session.get.return_value = resp_models

        resp_200 = MagicMock()
        resp_200.status_code = 200
        mock_session.post.side_effect = [resp_200]

        cliente = LlamaClient()
        cliente._make_request({"num_ctx": 8192})

        self.assertIsNone(cliente._num_ctx_respeitado)
        payload_enviado = mock_session.post.call_args.kwargs["json"]
        self.assertIn("num_ctx", payload_enviado)

    @patch('backend.core.llama_client.requests.Session')
    def test_flag_false_omite_num_ctx_da_proxima_chamada(self, mock_session_class):
        from backend.core.llama_client import LlamaClient

        mock_session = MagicMock()
        mock_session_class.return_value = mock_session
        resp_models = MagicMock()
        resp_models.status_code = 200
        mock_session.get.return_value = resp_models

        resp_200 = MagicMock()
        resp_200.status_code = 200
        mock_session.post.side_effect = [resp_200]

        cliente = LlamaClient()
        cliente._num_ctx_respeitado = False
        cliente._make_request({"num_ctx": 8192})

        payload_enviado = mock_session.post.call_args.kwargs["json"]
        self.assertNotIn("num_ctx", payload_enviado)
        self.assertEqual(mock_session.post.call_count, 1)  # sem retry


class TestPreCheckContexto(unittest.TestCase):
    """Testa o pre-check de contexto e o timeout por chamada no MariaRunner."""

    def _task(self, mensagem="Olá"):
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask
        return MariaTask(999, "PreCheck", "Teste", mensagem)

    def test_prompt_gigante_bloqueado_sem_retry(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner

        class ClienteConta:
            def __init__(self):
                self.model = "m"
                self.chamadas = 0

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                self.chamadas += 1
                return ("ok", None, 1, 1.0, 0.1)

        cliente = ClienteConta()
        runner = MariaRunner(cliente=cliente, ctx_size=100)
        resultado = runner.run(self._task("a" * 2000))

        self.assertFalse(resultado.contexto_ok)
        self.assertEqual(cliente.chamadas, 0)  # nunca chegou a enviar
        self.assertTrue(
            any(e.get("kind") == "LlamaClientError" for e in resultado.errors)
        )

    def test_prompt_normal_e_enviado(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner

        class ClienteNormal:
            def __init__(self):
                self.model = "m"
                self.chamadas = 0

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                self.chamadas += 1
                return ("Resposta em português.", None, 10, 5.0, 1.0)

        cliente = ClienteNormal()
        runner = MariaRunner(cliente=cliente)  # ctx default = LLAMA_NUM_CTX
        resultado = runner.run(self._task())

        self.assertTrue(resultado.contexto_ok)
        self.assertEqual(cliente.chamadas, 1)

    def test_callback_de_continuacao_usa_timeout_por_chamada(self):
        from backend.benchmarks.maria_bench.runners import maria_runner as modulo
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner

        class ClienteLeitura:
            def __init__(self):
                self.model = "m"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return ("", {"name": "listar_arquivos", "arguments": {}}, 5, 2.0, 0.5)

        capturado = {}

        def fake_encadear(cliente, historico, tool_call, schema, apos_cada_chamada=None, apos_cada_leitura=None):
            capturado["callback"] = apos_cada_chamada
            yield ("conteudo", None)

        runner = MariaRunner(cliente=ClienteLeitura())
        with patch.object(modulo, "encadear_leitura_stream", fake_encadear):
            runner.run(self._task())

        callback = capturado["callback"]
        limite = modulo.BENCHMARK_TIMEOUT_POR_CHAMADA
        with self.assertRaises(TimeoutError):
            callback(limite + 1, 5)
        # Abaixo do limite: não levanta (soma tokens via nonlocal internamente).
        callback(1, 7)


class TestCriarPlanilhaComLinhas(unittest.TestCase):
    """Testa criar_planilha_real e editar_planilha_real com o parâmetro linhas."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["PASTA_ARQUIVOS_GERADOS"] = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()
        os.environ.pop("PASTA_ARQUIVOS_GERADOS", None)

    def test_criar_sem_linhas_retrocompativel(self):
        """Chamada sem linhas mantém comportamento anterior (só cabeçalho)."""
        from backend.core.excel_handler import criar_planilha_real
        caminho = criar_planilha_real("teste", ["Nome", "Valor"])
        self.assertTrue(os.path.exists(caminho))
        df = pd.read_excel(caminho)
        self.assertEqual(list(df.columns), ["Nome", "Valor"])
        self.assertEqual(len(df), 0)

    def test_criar_com_linhas(self):
        """Cria planilha com dados e verifica conteúdo."""
        from backend.core.excel_handler import criar_planilha_real
        linhas = [{"Nome": "Ana", "Valor": 100}, {"Nome": "Bruno", "Valor": 200}]
        caminho = criar_planilha_real("teste", ["Nome", "Valor"], linhas=linhas)
        df = pd.read_excel(caminho)
        self.assertEqual(len(df), 2)
        self.assertEqual(df.iloc[0]["Nome"], "Ana")
        self.assertEqual(df.iloc[1]["Valor"], 200)

    def test_criar_descricao_ignorada_com_linhas(self):
        """v4.2.3: 'descricao' é ignorada (com warning) quando 'linhas' é
        fornecido — o arquivo sai com cabeçalho na linha 1, sem título."""
        from backend.core.excel_handler import criar_planilha_real
        from openpyxl import load_workbook
        linhas = [{"Nome": "Ana", "Valor": 1}]
        with self.assertLogs("backend.infrastructure.tools.excel_handler", level="WARNING") as logs:
            caminho = criar_planilha_real(
                "descricao_ignorada", ["Nome", "Valor"],
                descricao="Propósito da planilha",
                linhas=linhas,
            )
        self.assertTrue(any("descricao" in mensagem.lower() for mensagem in logs.output))
        wb = load_workbook(caminho)
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "Nome")  # cabeçalho, não descrição
        self.assertEqual(ws.cell(row=2, column=1).value, "Ana")
        wb.close()
        df = pd.read_excel(caminho)
        self.assertEqual(len(df), 1)

    def test_criar_coluna_ausente_fica_vazia(self):
        """Linha sem uma coluna deixa célula vazia."""
        from backend.core.excel_handler import criar_planilha_real
        linhas = [{"Nome": "Ana"}]  # "Valor" ausente
        caminho = criar_planilha_real("teste", ["Nome", "Valor"], linhas=linhas)
        df = pd.read_excel(caminho)
        self.assertEqual(df.iloc[0]["Nome"], "Ana")
        self.assertTrue(pd.isna(df.iloc[0]["Valor"]) or df.iloc[0]["Valor"] == "")

    def test_criar_chave_extra_ignorada(self):
        """Chave extra na linha não gera erro e é ignorada."""
        from backend.core.excel_handler import criar_planilha_real
        linhas = [{"Nome": "Ana", "Valor": 100, "Ignorar": "x"}]
        caminho = criar_planilha_real("teste", ["Nome", "Valor"], linhas=linhas)
        df = pd.read_excel(caminho)
        self.assertNotIn("Ignorar", df.columns)

    def test_editar_sem_linhas_retrocompativel(self):
        """editar_planilha sem linhas mantém comportamento anterior."""
        from backend.core.excel_handler import criar_planilha_real, editar_planilha_real
        caminho = criar_planilha_real("editar_teste", ["A"])
        nome = os.path.splitext(os.path.basename(caminho))[0]
        resultado = editar_planilha_real(nome, ["X", "Y"])
        df = pd.read_excel(resultado)
        self.assertEqual(list(df.columns), ["X", "Y"])
        self.assertEqual(len(df), 0)

    def test_editar_com_linhas(self):
        """editar_planilha com linhas escreve dados corretamente."""
        from backend.core.excel_handler import criar_planilha_real, editar_planilha_real
        caminho = criar_planilha_real("editar_dados", ["Nome"])
        nome = os.path.splitext(os.path.basename(caminho))[0]
        linhas = [{"Nome": "Carlos", "Valor": 50}]
        resultado = editar_planilha_real(nome, ["Nome", "Valor"], linhas=linhas)
        df = pd.read_excel(resultado)
        self.assertEqual(df.iloc[0]["Nome"], "Carlos")

    def test_limite_linhas_aplicado(self):
        """Número de linhas acima do limite é truncado silenciosamente."""
        from backend.core.excel_handler import criar_planilha_real
        from backend.core.config import get_max_linhas_por_chamada
        limite = get_max_linhas_por_chamada()
        linhas = [{"Col": i} for i in range(limite + 20)]
        caminho = criar_planilha_real("limite", ["Col"], linhas=linhas)
        df = pd.read_excel(caminho)
        self.assertLessEqual(len(df), limite)


class TestExtrairDadosPlanilha(unittest.TestCase):
    """Testa extrair_dados_planilha_real e a integração via executar_ferramenta_leitura."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["PASTA_ARQUIVOS_GERADOS"] = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()
        os.environ.pop("PASTA_ARQUIVOS_GERADOS", None)

    def test_arquivo_inexistente_levanta_value_error(self):
        from backend.core.excel_handler import extrair_dados_planilha_real
        with self.assertRaises(ValueError):
            extrair_dados_planilha_real("nao_existe")

    def test_planilha_vazia_retorna_zero_linhas(self):
        from backend.core.excel_handler import criar_planilha_real, extrair_dados_planilha_real
        caminho = criar_planilha_real("vazia", ["A", "B"])
        nome = os.path.splitext(os.path.basename(caminho))[0]
        resultado = extrair_dados_planilha_real(nome)
        self.assertEqual(resultado["colunas"], ["A", "B"])
        self.assertEqual(resultado["total_linhas"], 0)
        self.assertEqual(resultado["linhas"], [])
        self.assertFalse(resultado["tem_mais"])
        self.assertEqual(resultado["proximo_offset"], 0)

    def test_paginacao_basica(self):
        from unittest.mock import patch
        from backend.core.excel_handler import criar_planilha_real, extrair_dados_planilha_real
        # Limites separados: criação preserva todas as linhas (50) e a extração
        # pagina de fato (10 por página). Sem o patch, criar_planilha_real
        # truncaria as 15 linhas na criação (default 50) e a paginação nunca
        # veria tem_mais=True.
        limite = 10
        with patch("backend.infrastructure.tools.excel_handler.get_max_linhas_por_chamada", return_value=50), \
             patch("backend.infrastructure.tools.excel_handler.get_max_linhas_extracao", return_value=limite):
            linhas = [{"Col": i} for i in range(limite + 5)]
            caminho = criar_planilha_real("paginada", ["Col"], linhas=linhas)
            nome = os.path.splitext(os.path.basename(caminho))[0]

            pagina1 = extrair_dados_planilha_real(nome, offset=0)
            self.assertEqual(len(pagina1["linhas"]), limite)
            self.assertTrue(pagina1["tem_mais"])
            self.assertEqual(pagina1["proximo_offset"], limite)
            self.assertEqual(pagina1["total_linhas"], limite + 5)

            pagina2 = extrair_dados_planilha_real(nome, offset=pagina1["proximo_offset"])
            self.assertEqual(len(pagina2["linhas"]), 5)
            self.assertFalse(pagina2["tem_mais"])

    def test_offset_alem_do_total(self):
        from backend.core.excel_handler import criar_planilha_real, extrair_dados_planilha_real
        caminho = criar_planilha_real("curta", ["Col"], linhas=[{"Col": 1}, {"Col": 2}])
        nome = os.path.splitext(os.path.basename(caminho))[0]
        resultado = extrair_dados_planilha_real(nome, offset=100)
        self.assertEqual(resultado["linhas"], [])
        self.assertFalse(resultado["tem_mais"])
        self.assertEqual(resultado["proximo_offset"], 100)

    def test_planilha_com_descricao_detecta_cabecalho(self):
        from backend.core.excel_handler import criar_planilha_real, editar_planilha_real, extrair_dados_planilha_real
        # v4.2.3: criar_planilha_real ignora 'descricao' quando 'linhas' é
        # fornecido. O formato com descrição (cabeçalho na linha 3) é
        # exercitado via editar_planilha_real, que mantém descricao + linhas.
        caminho_criacao = criar_planilha_real("com_descricao", ["Nome", "Valor"])
        nome = os.path.splitext(os.path.basename(caminho_criacao))[0]
        editar_planilha_real(
            nome, ["Nome", "Valor"],
            descricao="Planilha de teste",
            linhas=[{"Nome": "Ana", "Valor": 10}],
        )
        resultado = extrair_dados_planilha_real(nome)
        self.assertEqual(resultado["colunas"], ["Nome", "Valor"])
        self.assertEqual(len(resultado["linhas"]), 1)
        self.assertEqual(resultado["linhas"][0]["Nome"], "Ana")

    def test_offset_negativo_tratado_como_zero(self):
        from backend.core.excel_handler import criar_planilha_real, extrair_dados_planilha_real
        caminho = criar_planilha_real("neg", ["Col"], linhas=[{"Col": 1}])
        nome = os.path.splitext(os.path.basename(caminho))[0]
        resultado = extrair_dados_planilha_real(nome, offset=-5)
        self.assertEqual(resultado["offset_atual"], 0)

    def test_executar_ferramenta_leitura_retorna_json_valido(self):
        import json
        from backend.core.excel_handler import criar_planilha_real
        from backend.core.tools_schema import executar_ferramenta_leitura
        caminho = criar_planilha_real("via_ferramenta", ["A"], linhas=[{"A": 1}])
        nome = os.path.splitext(os.path.basename(caminho))[0]
        resultado_str = executar_ferramenta_leitura("extrair_dados_planilha", {"nome_arquivo": nome})
        dados = json.loads(resultado_str)
        self.assertEqual(dados["colunas"], ["A"])
        self.assertEqual(dados["total_linhas"], 1)

    def test_ferramenta_registrada_em_ferramentas_leitura(self):
        from backend.core.tools_schema import FERRAMENTAS_LEITURA
        self.assertIn("extrair_dados_planilha", FERRAMENTAS_LEITURA)

    def test_linha_cabecalho_override(self):
        """v4.2.3: linha_cabecalho=2 (0-indexado) extrai cabeçalhos da linha 3
        do Excel (formato com descrição, gerado via editar_planilha_real)."""
        from backend.core.excel_handler import criar_planilha_real, editar_planilha_real, extrair_dados_planilha_real
        caminho_criacao = criar_planilha_real("override_cab", ["Nome", "Valor"])
        nome = os.path.splitext(os.path.basename(caminho_criacao))[0]
        editar_planilha_real(
            nome, ["Nome", "Valor"],
            descricao="Título que será ignorado pela leitura",
            linhas=[{"Nome": "Ana", "Valor": 7}],
        )
        resultado = extrair_dados_planilha_real(nome, linha_cabecalho=2)
        self.assertEqual(resultado["colunas"], ["Nome", "Valor"])
        self.assertEqual(len(resultado["linhas"]), 1)
        self.assertEqual(resultado["linhas"][0]["Nome"], "Ana")

    def test_limite_linhas_extracao(self):
        """v4.2.3: limite_linhas reduz o lote, mas o teto do modelo prevalece."""
        from backend.core.excel_handler import criar_planilha_real, extrair_dados_planilha_real
        with patch("backend.infrastructure.tools.excel_handler.get_max_linhas_por_chamada", return_value=50), \
             patch("backend.infrastructure.tools.excel_handler.get_max_linhas_extracao", return_value=10):
            linhas = [{"Col": i} for i in range(15)]
            caminho = criar_planilha_real("limitada", ["Col"], linhas=linhas)
            nome = os.path.splitext(os.path.basename(caminho))[0]

            pagina = extrair_dados_planilha_real(nome, limite_linhas=4)
            self.assertEqual(len(pagina["linhas"]), 4)
            self.assertTrue(pagina["tem_mais"])
            self.assertEqual(pagina["proximo_offset"], 4)

            # limite_linhas acima do teto do modelo → teto (10) prevalece
            pagina_teto = extrair_dados_planilha_real(nome, limite_linhas=100)
            self.assertEqual(len(pagina_teto["linhas"]), 10)
            self.assertTrue(pagina_teto["tem_mais"])

            # sem limite_linhas → comportamento anterior (teto do modelo)
            pagina_padrao = extrair_dados_planilha_real(nome)
            self.assertEqual(len(pagina_padrao["linhas"]), 10)

    def test_tipos_incluidos_no_json(self):
        """v4.2.3: o retorno inclui 'tipos' com o dtype pandas de cada coluna."""
        from backend.core.excel_handler import criar_planilha_real, extrair_dados_planilha_real
        caminho = criar_planilha_real(
            "com_tipos", ["Nome", "Idade"],
            linhas=[{"Nome": "Ana", "Idade": 30}],
        )
        nome = os.path.splitext(os.path.basename(caminho))[0]
        resultado = extrair_dados_planilha_real(nome)
        self.assertIn("tipos", resultado)
        self.assertIsInstance(resultado["tipos"], dict)
        self.assertEqual(set(resultado["tipos"].keys()), {"Nome", "Idade"})
        for tipo in resultado["tipos"].values():
            self.assertIsInstance(tipo, str)
        # Detecção automática de cabeçalho segue funcionando junto com 'tipos'
        self.assertEqual(resultado["colunas"], ["Nome", "Idade"])


class TestTarefa26Traducao(unittest.TestCase):
    """v4.2.2: Task 26 do benchmark — tradução mandarim → inglês de planilha
    real (produtos). Valida o desenho da tarefa e a cópia da fixture real
    (backend/benchmarks/maria_bench/fixtures/produtos_mandarim.xlsx)."""

    def test_estrutura_da_task_26(self):
        from backend.benchmarks.maria_bench.tasks import load_all_maria_tasks

        tarefas = {t.id: t for t in load_all_maria_tasks()}
        self.assertIn(26, tarefas)
        task = tarefas[26]
        self.assertEqual(task.expected_tool, "criar_planilha")
        self.assertEqual(
            task.tools_obrigatorios,
            ["extrair_dados_planilha", "criar_planilha"],
        )
        self.assertEqual(task.fixtures, ["produtos_mandarim"])
        self.assertEqual(task.confirm_sequence, ["sim"])
        self.assertEqual(
            task.expected_args_subset,
            {"nome_arquivo": "produtos_traduzidos"},
        )
        # FIX-1/FIX-2/FIX-3 (diagnóstico 2026-09-08): a user_message não pode
        # sugerir edição ("preencha a coluna"/"edite") e o context não pode
        # pré-declarar o arquivo disponível — o modelo deve LÊ-LO via
        # extrair_dados_planilha antes de criar a planilha nova.
        msg_lower = task.user_message.lower()
        self.assertNotIn("preencha a coluna", msg_lower)
        self.assertNotIn("edite", msg_lower)
        self.assertTrue(
            "crie" in msg_lower or "nova planilha" in msg_lower,
            "user_message deve instruir a criação de arquivo novo",
        )
        self.assertEqual(task.context, [])

    def test_fixture_produtos_mandarim_copiada(self):
        """A fixture produtos_mandarim.xlsx é copiada do diretório real de
        fixtures (backend/benchmarks/maria_bench/fixtures/), não gerada programaticamente."""
        import pandas as pd
        from unittest.mock import patch
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner

        task = MariaTask(
            id=26,
            name="Tradução de planilha",
            description="desc",
            user_message="Preencha a columna english description da planilha produtos.",
            fixtures=["produtos_mandarim"],
            category=MariaTaskCategory.CRIAR_PLANILHA,
        )
        with tempfile.TemporaryDirectory() as tmp:
            with patch(
                "backend.benchmarks.maria_bench.runners.maria_runner.BENCHMARK_ARQUIVOS_DIR", tmp
            ):
                MariaRunner._garantir_planilha_existente(task)
                caminho = os.path.join(tmp, "produtos_mandarim.xlsx")
                self.assertTrue(os.path.exists(caminho))
                df = pd.read_excel(caminho)
                self.assertEqual(
                    list(df.columns),
                    ["model", "product", "english description", "NCM"],
                )
                self.assertEqual(len(df), 6)
                self.assertEqual(df.iloc[0]["model"], "QFY000013")
                # Rodar de novo não duplica (idempotente: arquivo existe)
                MariaRunner._garantir_planilha_existente(task)
                df2 = pd.read_excel(caminho)
                self.assertEqual(len(df2), 6)


class TestValidacaoDadosArquivoGerado(unittest.TestCase):
    """Item A: valida coluna obrigatória preenchida no .xlsx gerado (Task 26).

    Cobre a função `_validar_coluna_preenchida` (casos diretos) e a integração
    com o MariaRunner.run() — arquivo gerado só com cabeçalho agora resulta em
    dados_arquivo_validos=False + erro "DadosIncompletos".
    """

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def _escrever_xlsx(self, nome: str, dados: list[dict]) -> str:
        caminho = os.path.join(self.tmp.name, nome)
        pd.DataFrame(dados).to_excel(caminho, index=False)
        return caminho

    def test_coluna_preenchida_em_todas_as_linhas(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import _validar_coluna_preenchida
        caminho = self._escrever_xlsx("preenchida.xlsx", [
            {"english description": "Foo"},
            {"english description": "Bar"},
            {"english description": "Baz"},
        ])
        self.assertEqual(
            _validar_coluna_preenchida(caminho, "english description"), (True, "")
        )

    def test_coluna_vazia_em_pelo_menos_uma_linha(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import _validar_coluna_preenchida
        caminho = self._escrever_xlsx("com_vazia.xlsx", [
            {"english description": "Foo"},
            {"english description": None},
            {"english description": "Baz"},
        ])
        valido, motivo = _validar_coluna_preenchida(caminho, "english description")
        self.assertFalse(valido)
        self.assertIn("1/3", motivo)

    def test_coluna_inexistente_no_arquivo(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import _validar_coluna_preenchida
        caminho = self._escrever_xlsx("sem_coluna.xlsx", [{"model": "QFY000013"}])
        valido, motivo = _validar_coluna_preenchida(caminho, "english description")
        self.assertFalse(valido)
        self.assertIn("english description", motivo)

    def test_arquivo_inexistente_retorna_false_sem_excecao(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import _validar_coluna_preenchida
        caminho = os.path.join(self.tmp.name, "nao_existe.xlsx")
        valido, motivo = _validar_coluna_preenchida(caminho, "english description")
        self.assertFalse(valido)
        self.assertIn("não encontrado", motivo)

    def test_case_insensitive_na_coluna(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import _validar_coluna_preenchida
        caminho = self._escrever_xlsx("case.xlsx", [
            {"English Description": "Foo"},
            {"English Description": "Bar"},
        ])
        self.assertEqual(
            _validar_coluna_preenchida(caminho, "english description"), (True, "")
        )

    def test_task_sem_coluna_dados_obrigatoria_nao_valida_arquivo(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteTexto:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                return ("Olá! Posso ajudar com planilhas e documentos.", None, 5, 2.0, 1.0)

        task = MariaTask(
            1, "Conversa simples", "Saudação sem ferramenta",
            "Olá, como você pode me ajudar?",
            category=MariaTaskCategory.CONVERSA,
        )
        runner = MariaRunner(cliente=ClienteTexto())
        with patch("backend.benchmarks.maria_bench.runners.maria_runner.BENCHMARK_ARQUIVOS_DIR", self.tmp.name):
            with patch(
                "backend.benchmarks.maria_bench.runners.maria_runner._validar_coluna_preenchida",
                return_value=(True, ""),
            ) as mock_validar:
                resultado = runner.run(task)
        self.assertTrue(resultado.dados_arquivo_validos)
        mock_validar.assert_not_called()

    def test_runner_task_26_arquivo_somente_cabecalho_marca_dados_invalidos(self):
        from backend.benchmarks.maria_bench.runners.maria_runner import MariaRunner
        from backend.benchmarks.maria_bench.tasks.task_schema import MariaTask, MariaTaskCategory

        class ClienteCriaPlanilha:
            model = "modelo-teste"

            def chat_com_tools_stream_com_metricas(self, **kwargs):
                # Modo de falha 1 do relatório: modelo manda as colunas certas
                # sem `linhas` → arquivo é criado só com cabeçalho.
                return ("", {
                    "name": "criar_planilha",
                    "arguments": {
                        "nome_arquivo": "produtos_traduzidos",
                        "colunas": ["model", "product", "english description", "NCM"],
                    },
                }, 5, 2.0, 1.0)

        task = MariaTask(
            26, "Tradução de planilha (Mandarim → Inglês)", "desc",
            "Preencha a coluna english description da planilha produtos.",
            expected_tool="criar_planilha",
            confirm_sequence=["sim"],
            category=MariaTaskCategory.CRIAR_PLANILHA,
            coluna_dados_obrigatoria="english description",
        )
        runner = MariaRunner(cliente=ClienteCriaPlanilha())
        with tempfile.TemporaryDirectory() as tmp:
            with patch("backend.benchmarks.maria_bench.runners.maria_runner.BENCHMARK_ARQUIVOS_DIR", tmp):
                resultado = runner.run(task)
        self.assertFalse(resultado.dados_arquivo_validos)
        kinds = [e["kind"] for e in resultado.errors]
        self.assertIn("DadosIncompletos", kinds)
        self.assertIn("linha", resultado.errors[0]["message"])
        self.assertFalse(resultado.runtime_ok)


if __name__ == "__main__":
    unittest.main()
